#!/usr/bin/env python3
"""
Sobeys receipt extractor with canonical DB write support.

Flow mirrors Costco/Instacart loaders:
- parse receipt PDFs
- classify item categories using shared logic (+ optional AI fallback)
- write XLSX outputs (items, receipts, category totals, notes)
- optionally upsert canonical records into budget.expenses + budget.expense_items
"""

from __future__ import annotations

import argparse
import difflib
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from budget_category_ai import AICategoryEngine
from budget_category_logic import CATEGORY_ORDER, DEFAULT_CATEGORY, deterministic_category, normalize_for_match

PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = PROJECT_DIR / "receipts" / "Sobeys"
DEFAULT_OUT_XLSX = PROJECT_DIR / "sobeys_receipt_items_categorized.xlsx"
AI_SUGGESTIONS_PATH = PROJECT_DIR / "ai_category_cache.json"
SUPPORTED_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".heic", ".heif"}

AMOUNT_RE = re.compile(r"(\$?-?\d+\.\d{2}-?)\s*$")
MONEY_TOKEN_RE = re.compile(r"\$?-?\d+\.\d{2}-?")
MONEY_STRICT_RE = re.compile(r"\$-?\d+\.\d{2}-?")
WEIGHT_DETAIL_RE = re.compile(
    r"(?P<qty>\d+(?:\.\d+)?)\s*(?P<unit>kg|g|lb|lbs)\s*@\s*\$?(?P<unit_cost>\d+\.\d{2})(?:\s*/\s*(?P<per_unit>kg|g|lb|lbs))?",
    re.I,
)
COUNT_DETAIL_RE = re.compile(
    r"(?P<qty>\d+(?:\.\d+)?)\s*@\s*(?P<bundle>\d+)\s*/\s*\$?(?P<unit_cost>\d+\.\d{2})",
    re.I,
)
DATE_PREFIX_RE = re.compile(r"^(\d{8})")
FILENAME_AMOUNT_RE = re.compile(r"^\d{8}_sobeys_([0-9]+(?:[._][0-9]{1,2})?)$", re.I)

SKIP_PATTERNS = [
    r"\bSUBTOTAL\b",
    r"\bTOTAL\b",
    r"\bTAX\b",
    r"\bGST\b",
    r"\bPST\b",
    r"\bHST\b",
    r"\bBALANCE\b",
    r"\bCHANGE\b",
    r"\bTENDER\b",
    r"\bTENDER\s+CASH\b",
    r"\bCASH\b",
    r"\bDEBIT\b",
    r"\bCREDIT\b",
    r"\bVISA\b",
    r"\bMASTERCARD\b",
    r"\bTHANK\b",
    r"\bSOBEYS\b",
    r"\bSCENE\b",
    r"\bPOINTS\b",
    r"\bSTORE\b",
    r"\bCASHIER\b",
    r"\bTRANSACTION\b",
    r"\bAPPROVED\b",
    r"\bAUTH\b",
    r"\bYOU\s*SAVED\b",
    r"\bINSTANT\s*SAVINGS?\b",
    r"\bDISCOUNT\b",
    r"\bSCENE\+?\s*POINTS?\b",
]
SKIP_RE = re.compile("|".join(SKIP_PATTERNS), re.I)
SECTION_HEADER_WORDS = {
    "produce",
    "dairy",
    "bakery",
    "meat",
    "seafood",
    "frozen",
    "deli",
    "pantry",
    "beverages",
    "snacks",
    "household",
    "pharmacy",
}
NON_ITEM_TOKENS = {
    "grocery",
    "produce",
    "pro",
    "c",
    "bc",
    "bl",
    "ps",
}
SOBEYS_NAME_OVERRIDES = {
    "piglet foeycr en": "Apples Honeycrisp",
    "mango ale": "Mango Spears",
    "dill freeze ord 10g": "Dill Freeze Drd 10G",
    "cake mx carrotsprimst": "Cake Mix Carrot",
}
SOBEYS_NAME_CANDIDATES = [
    "Apples Honeycrisp",
    "Mango Spears",
    "Basil Freeze Drd 8G",
    "Dill Freeze Drd 10G",
    "Cottage Cheese 1%MF",
    "Sour Cream Light",
    "Topping Dessert 1L",
    "Asparagus",
    "Bananas",
    "Cabbage Green",
    "Onions Green",
    "Onions Red",
]

AI_ENGINE = AICategoryEngine(AI_SUGGESTIONS_PATH)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract Sobeys receipt data and optionally write canonical DB records.")
    parser.add_argument("--input-path", default=str(DEFAULT_INPUT), help="PDF file or directory containing Sobeys receipt PDFs.")
    parser.add_argument("--out-xlsx", default=str(DEFAULT_OUT_XLSX), help="Output Excel workbook path.")
    parser.add_argument("--write-db", action="store_true", help="Write parsed canonical expense records into Postgres.")
    parser.add_argument(
        "--db-dsn",
        default=os.environ.get("HOME_BUDGET_PG_DSN", ""),
        help=(
            "Postgres DSN for --write-db. Do not include password here. "
            "Use ~/.pgpass or PGPASSWORD/HOME_BUDGET_PGPASSWORD env vars. "
            "Default: env HOME_BUDGET_PG_DSN or dbname=home_budget."
        ),
    )
    parser.add_argument("--db-schema", default="budget", help="Target schema for --write-db (default: budget).")
    return parser.parse_args()


def normalize_line(line: str) -> str:
    return re.sub(r"\s+", " ", line).strip()


def parse_amount_token(token: str) -> float:
    token = token.replace("$", "")
    if token.endswith("-") and not token.startswith("-"):
        token = f"-{token[:-1]}"
    return float(token)


def parse_money_value(value: Optional[str]) -> Optional[float]:
    if not value:
        return None
    m = re.search(r"(?P<prefix>-)?\$?\s*(?P<num>[0-9]+(?:,[0-9]{3})*(?:\.[0-9]{2})?)(?P<suffix>-)?", value)
    if not m:
        return None
    out = float(m.group("num").replace(",", ""))
    if m.group("prefix") or m.group("suffix"):
        out = -out
    return out


def extract_inputs(input_path: Path) -> List[Path]:
    if input_path.is_file() and input_path.suffix.lower() in {".pdf", *SUPPORTED_IMAGE_EXTS}:
        return [input_path]
    if input_path.is_dir():
        files: List[Path] = []
        for p in input_path.rglob("*"):
            if p.is_file() and p.suffix.lower() in {".pdf", *SUPPORTED_IMAGE_EXTS}:
                files.append(p)
        return sorted(files)
    return []


def pdf_text(pdf_path: Path) -> str:
    parts: List[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    text = "\n".join(parts)
    if text.strip():
        return text
    return ocr_pdf_text(pdf_path)


def ocr_pdf_text(pdf_path: Path) -> str:
    tesseract_bin = shutil.which("tesseract")
    if not tesseract_bin:
        return ""
    try:
        import pypdfium2 as pdfium  # type: ignore
    except Exception:
        return ""

    ocr_parts: List[str] = []
    doc = None
    try:
        doc = pdfium.PdfDocument(str(pdf_path))
        for idx in range(len(doc)):
            page = doc[idx]
            # 300 DPI render target for better OCR reliability.
            pil_img = page.render(scale=300 / 72).to_pil()
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                tmp_path = Path(tmp.name)
            try:
                pil_img.save(tmp_path)
                proc = subprocess.run(
                    [tesseract_bin, str(tmp_path), "stdout", "--psm", "11"],
                    check=False,
                    capture_output=True,
                    text=True,
                )
                if proc.returncode == 0 and proc.stdout:
                    ocr_parts.append(proc.stdout)
            finally:
                try:
                    tmp_path.unlink(missing_ok=True)
                except Exception:
                    pass
    except Exception:
        return ""
    finally:
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass
    return "\n".join(ocr_parts)


def ocr_image_text(image_path: Path) -> str:
    tesseract_bin = shutil.which("tesseract")
    if not tesseract_bin:
        return ""

    def run_tesseract(path: Path) -> str:
        proc = subprocess.run(
            [tesseract_bin, str(path), "stdout", "--psm", "11"],
            check=False,
            capture_output=True,
            text=True,
        )
        if proc.returncode == 0 and proc.stdout:
            return proc.stdout
        return ""

    text = run_tesseract(image_path)
    if text.strip():
        return text

    if image_path.suffix.lower() not in {".heic", ".heif"}:
        return text

    # HEIC fallback: convert to PNG and OCR.
    heif_convert_bin = shutil.which("heif-convert")
    if not heif_convert_bin:
        return text
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp_png = Path(tmp.name)
    try:
        proc = subprocess.run(
            [heif_convert_bin, str(image_path), str(tmp_png)],
            check=False,
            capture_output=True,
            text=True,
        )
        if proc.returncode != 0:
            return text
        return run_tesseract(tmp_png)
    finally:
        try:
            tmp_png.unlink(missing_ok=True)
        except Exception:
            pass


def document_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return pdf_text(path)
    return ocr_image_text(path)


def document_text_alt(path: Path) -> str:
    # Alternate OCR pass used for line-repair when primary OCR text is noisy.
    if path.suffix.lower() != ".pdf":
        return ""
    tesseract_bin = shutil.which("tesseract")
    if not tesseract_bin:
        return ""
    try:
        import pypdfium2 as pdfium  # type: ignore
    except Exception:
        return ""
    out: List[str] = []
    doc = None
    try:
        doc = pdfium.PdfDocument(str(path))
        for idx in range(len(doc)):
            page = doc[idx]
            pil_img = page.render(scale=300 / 72).to_pil()
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                tmp_path = Path(tmp.name)
            try:
                pil_img.save(tmp_path)
                proc = subprocess.run(
                    [tesseract_bin, str(tmp_path), "stdout", "--psm", "4"],
                    check=False,
                    capture_output=True,
                    text=True,
                )
                if proc.returncode == 0 and proc.stdout:
                    out.append(proc.stdout)
            finally:
                try:
                    tmp_path.unlink(missing_ok=True)
                except Exception:
                    pass
    except Exception:
        return ""
    finally:
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass
    return "\n".join(out)


def parse_receipt_amounts(text: str) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    subtotal = None
    total = None
    tax_total = None
    for raw in text.splitlines():
        line = normalize_line(raw)
        m_sub = re.search(r"\bSUBTOTAL\b.*?(-?\d+\.\d{2})$", line, re.I)
        if m_sub:
            subtotal = float(m_sub.group(1))
            continue
        m_tax = re.search(r"\b(?:TOTAL\s+TAX|TAX)\b.*?(-?\d+\.\d{2})$", line, re.I)
        if m_tax:
            tax_total = float(m_tax.group(1))
            continue
        m_total = re.search(r"^\*?\s*TOTAL\b(?!\s*(?:TAX|ITEMS?)\b).*?(-?\d+\.\d{2})$", line, re.I)
        if m_total:
            total = float(m_total.group(1))
    return subtotal, tax_total, total


def likely_item_line(line: str) -> bool:
    if not line:
        return False
    if SKIP_RE.search(line):
        return False
    if not MONEY_TOKEN_RE.search(line):
        return False
    if not re.search(r"[A-Za-z]", line):
        return False
    return True


def parse_item_line(line: str) -> Optional[Tuple[str, float]]:
    line = normalize_line(line)
    if re.search(r"\bsaved\b", line, re.I):
        return None
    matches = list(MONEY_TOKEN_RE.finditer(line))
    if not matches:
        return None
    m = matches[-1]
    amount = parse_amount_token(m.group(0))
    desc = line[:m.start()].strip(" -:")
    desc = normalize_line(desc)
    if not desc:
        return None
    return desc, amount


def parse_amount_only_line(line: str) -> Optional[float]:
    line = normalize_line(line)
    if re.search(r"\bsaved\b|\binstant\s+savings?\b", line, re.I):
        return None
    if "@" in line and re.search(r"\b(?:kg|g|lb|lbs)\b", line, re.I):
        return None
    matches = list(MONEY_STRICT_RE.finditer(line))
    if not matches:
        return None
    return parse_amount_token(matches[-1].group(0))


def parse_weight_detail_line(line: str) -> Optional[Tuple[float, str, float]]:
    ln = normalize_line(line)
    m = WEIGHT_DETAIL_RE.search(ln)
    if not m:
        return None
    try:
        qty = float(m.group("qty"))
        unit = m.group("unit").lower()
        unit_cost = float(m.group("unit_cost"))
        return qty, unit, unit_cost
    except Exception:
        return None


def parse_count_detail_line(line: str) -> Optional[Tuple[float, float]]:
    ln = normalize_line(line)
    m = COUNT_DETAIL_RE.search(ln)
    if not m:
        return None
    try:
        qty = float(m.group("qty"))
        unit_cost = float(m.group("unit_cost"))
        return qty, unit_cost
    except Exception:
        return None


def find_item_details(
    raw_lines: List[str],
    start_line_1based: int,
) -> Tuple[Optional[float], Optional[str], Optional[float], Optional[float], Optional[int]]:
    # Returns: weight_qty, weight_unit, weight_unit_cost, count_qty, detail_line_no
    for j in range(start_line_1based, min(len(raw_lines), start_line_1based + 5) + 1):
        line = raw_lines[j - 1]
        window_parts = [line]
        for k in range(j + 1, min(len(raw_lines), j + 4) + 1):
            nxt = raw_lines[k - 1]
            if nxt:
                window_parts.append(nxt)
        window_line = normalize_line(" ".join(window_parts))

        parsed_weight = parse_weight_detail_line(window_line)
        if parsed_weight is not None:
            qty, unit, unit_cost = parsed_weight
            return qty, unit, unit_cost, None, j
        parsed_count = parse_count_detail_line(window_line)
        if parsed_count is not None:
            qty, unit_cost = parsed_count
            return None, None, unit_cost, qty, j
    return None, None, None, None, None


def looks_desc_candidate(line: str) -> bool:
    ln = normalize_line(line)
    if not ln:
        return False
    if SKIP_RE.search(ln):
        return False
    if MONEY_TOKEN_RE.search(ln):
        return False
    if not re.search(r"[A-Za-z]", ln):
        return False
    if ln.lower() in NON_ITEM_TOKENS:
        return False
    alpha_only = re.sub(r"[^A-Za-z]+", "", ln)
    if len(alpha_only) <= 2:
        return False
    # Ignore uppercase section headers like "PRODUCE".
    words = re.findall(r"[A-Za-z]+", ln)
    if words:
        joined = " ".join(w.lower() for w in words)
        if joined in SECTION_HEADER_WORDS:
            return False
        if ln == ln.upper() and len(words) <= 3 and all(w.lower() in SECTION_HEADER_WORDS for w in words):
            return False
    # Avoid obvious headers/sections.
    if len(ln) > 80:
        return False
    return True


def looks_non_item_desc(desc: str) -> bool:
    d = normalize_line(desc).lower()
    if not d:
        return True
    blocked = [
        r"^you s[a-z]{2,}",
        r"\binstant savings?\b",
        r"\bdiscounts?\b",
        r"\bspecials?\b",
        r"\bsub\s*total\b",
        r"\bsbtotal\b",
        r"\btotal\b",
        r"\btax\b",
        r"\bgst\b",
        r"\bpst\b",
        r"\bhst\b",
        r"\bbalance\b",
        r"\bchange\b",
        r"\bpurchase\b",
        r"\bpoints?\b",
        r"\bscene\b",
        r"^\s*/\s*kg\s*$",
        r"^\+?eh\b",
    ]
    if any(re.search(p, d) for p in blocked):
        return True
    if d in NON_ITEM_TOKENS:
        return True
    words = re.findall(r"[A-Za-z]+", d)
    if words:
        if " ".join(words) in SECTION_HEADER_WORDS:
            return True
        if len(words) <= 3 and all(w in SECTION_HEADER_WORDS for w in words):
            return True
    alnum = [c for c in d if c.isalnum()]
    if alnum:
        digit_ratio = sum(c.isdigit() for c in alnum) / len(alnum)
        if digit_ratio > 0.45:
            return True
    alpha_only = re.sub(r"[^a-z]+", "", d)
    if len(alpha_only) <= 2:
        return True
    return False


def canonicalize_ocr_item_name(desc: str) -> str:
    raw = normalize_line(desc)
    key = normalize_for_match(raw)
    exact = SOBEYS_NAME_OVERRIDES.get(key)
    if exact:
        return exact

    # Fuzzy normalization against known Sobeys candidates.
    best_name = None
    best_score = 0.0
    for candidate in SOBEYS_NAME_CANDIDATES:
        score = difflib.SequenceMatcher(None, key, normalize_for_match(candidate)).ratio()
        if score > best_score:
            best_score = score
            best_name = candidate
    if best_name and best_score >= 0.72:
        return best_name
    return raw


def classify_category_with_source(desc: str, allow_ai: bool = True) -> Tuple[str, str]:
    runtime_default = DEFAULT_CATEGORY if DEFAULT_CATEGORY in CATEGORY_ORDER else (CATEGORY_ORDER[0] if CATEGORY_ORDER else DEFAULT_CATEGORY)
    deterministic, source = deterministic_category(desc)
    if deterministic not in CATEGORY_ORDER:
        deterministic = runtime_default
        source = "default"
    if source != "default":
        return deterministic, source

    if allow_ai:
        cache_key = normalize_for_match(desc)
        cached = AI_ENGINE.get_cached_category_by_key(cache_key)
        if cached in CATEGORY_ORDER:
            return cached, "ai_cache"
        ai_category = AI_ENGINE.try_single_category(desc, source_label="Sobeys item")
        if ai_category in CATEGORY_ORDER:
            return ai_category, "ai"
    return deterministic, "default"


def _desc_quality(desc: str) -> float:
    d = normalize_line(desc)
    if not d:
        return 0.0
    alpha = re.sub(r"[^A-Za-z]+", "", d)
    words = [w for w in re.findall(r"[A-Za-z]+", d)]
    score = 0.0
    score += min(len(alpha), 24) / 24.0
    score += min(len(words), 5) / 5.0
    if any(ch.isdigit() for ch in d):
        score -= 0.25
    if looks_non_item_desc(d):
        score -= 0.5
    return score


def _build_amount_desc_candidates(text: str) -> Dict[float, List[str]]:
    by_amount: Dict[float, List[str]] = {}
    for raw in text.splitlines():
        parsed = parse_item_line(raw)
        if not parsed:
            continue
        desc, amount = parsed
        if looks_non_item_desc(desc):
            continue
        by_amount.setdefault(round(amount, 2), []).append(desc)
    return by_amount


def _repair_desc_from_alt(desc: str, amount: float, alt_candidates: Dict[float, List[str]]) -> str:
    key = round(float(amount), 2)
    candidates = alt_candidates.get(key, [])
    if not candidates:
        return desc

    def noise_flags(text: str) -> Tuple[int, int]:
        words = re.findall(r"[A-Za-z0-9]+", normalize_line(text))
        short_words = sum(1 for w in words if len(w) <= 2)
        digit_tokens = sum(1 for w in words if any(ch.isdigit() for ch in w))
        return short_words, digit_tokens

    best = desc
    best_score = _desc_quality(desc)
    base_short, base_digit = noise_flags(desc)
    for cand in candidates:
        s = _desc_quality(cand)
        cand_short, cand_digit = noise_flags(cand)
        # Prefer clearly cleaner alternate OCR when base text looks noisy.
        cleaner = (cand_digit < base_digit) or (cand_short + 1 < base_short)
        if cleaner and s >= best_score - 0.25:
            best = cand
            best_score = s
            base_short, base_digit = cand_short, cand_digit
            continue
        if s > best_score + 0.20:
            best = cand
            best_score = s
            base_short, base_digit = cand_short, cand_digit
    return best


def parse_items(receipt_id: str, text: str, alt_text: str = "") -> List[Dict]:
    out: List[Dict] = []
    alt_candidates = _build_amount_desc_candidates(alt_text) if alt_text else {}
    raw_lines = [normalize_line(x) for x in text.splitlines()]
    idx = 1
    while idx <= len(raw_lines):
        line = raw_lines[idx - 1]
        parsed = parse_item_line(line)
        if parsed is not None:
            desc, amount = parsed
            desc = canonicalize_ocr_item_name(desc)
            if alt_candidates:
                desc = _repair_desc_from_alt(desc, amount, alt_candidates)
            if not looks_non_item_desc(desc):
                weight_qty, weight_unit, detail_unit_cost, count_qty, _ = find_item_details(raw_lines, idx)
                category, source = classify_category_with_source(desc, allow_ai=True)
                out.append(
                    {
                        "receipt_id": receipt_id,
                        "line_no": idx,
                        "item": desc,
                        "amount_cad": amount,
                        "unit_qty": count_qty if count_qty is not None else (None if weight_qty is not None else 1.0),
                        "unit_cost": detail_unit_cost if detail_unit_cost is not None else amount,
                        "weight_qty": weight_qty,
                        "weight_unit": weight_unit,
                        "line_total": amount,
                        "budget_category": category,
                        "category_source": source,
                        "raw_line": line,
                    }
                )
                idx += 1
                continue

        # Multi-line OCR pattern: description line followed by amount-only line
        # with optional blank/noisy lines in-between.
        if looks_desc_candidate(line):
            used_j = None
            amount = None
            raw_amount_line = None
            for j in range(idx + 1, min(len(raw_lines), idx + 10) + 1):
                candidate = raw_lines[j - 1]
                if not candidate:
                    continue
                # Stop if the next item description starts before we found a price.
                if looks_desc_candidate(candidate):
                    break
                amt = parse_amount_only_line(candidate)
                if amt is None:
                    continue
                amount = amt
                used_j = j
                raw_amount_line = candidate
                break
            if amount is not None and used_j is not None and raw_amount_line is not None:
                desc = canonicalize_ocr_item_name(line)
                if alt_candidates:
                    desc = _repair_desc_from_alt(desc, amount, alt_candidates)
                if not looks_non_item_desc(desc):
                    weight_qty, weight_unit, detail_unit_cost, count_qty, _ = find_item_details(raw_lines, used_j)
                    category, source = classify_category_with_source(desc, allow_ai=True)
                    out.append(
                        {
                            "receipt_id": receipt_id,
                            "line_no": idx,
                            "item": desc,
                            "amount_cad": amount,
                            "unit_qty": count_qty if count_qty is not None else (None if weight_qty is not None else 1.0),
                            "unit_cost": detail_unit_cost if detail_unit_cost is not None else amount,
                            "weight_qty": weight_qty,
                            "weight_unit": weight_unit,
                            "line_total": amount,
                            "budget_category": category,
                            "category_source": source,
                            "raw_line": f"{line} | {raw_amount_line}",
                        }
                    )
                    idx = used_j + 1
                    continue
        idx += 1
    return remove_likely_ocr_repeats(out)


def remove_likely_ocr_repeats(items: List[Dict]) -> List[Dict]:
    # Some OCR runs duplicate large receipt sections later in the text stream.
    # Keep the first occurrence and drop identical rows that reappear far away.
    first_seen_line: Dict[Tuple[str, float, Optional[float], Optional[float], Optional[str]], int] = {}
    filtered: List[Dict] = []
    for row in items:
        key = (
            normalize_for_match(str(row.get("item") or "")),
            round(float(row.get("line_total") or 0.0), 2),
            round(float(row.get("unit_cost")), 2) if row.get("unit_cost") is not None else None,
            round(float(row.get("weight_qty")), 3) if row.get("weight_qty") is not None else None,
            str(row.get("weight_unit") or "").lower() or None,
        )
        line_no = int(row.get("line_no") or 0)
        first_line = first_seen_line.get(key)
        if first_line is None:
            first_seen_line[key] = line_no
            filtered.append(row)
            continue
        # Only treat as OCR repeat when the same row appears much later.
        if line_no - first_line > 150:
            continue
        filtered.append(row)
    return filtered


def fallback_category() -> str:
    if "Cash/Unknown" in CATEGORY_ORDER:
        return "Cash/Unknown"
    if DEFAULT_CATEGORY in CATEGORY_ORDER:
        return DEFAULT_CATEGORY
    return CATEGORY_ORDER[0] if CATEGORY_ORDER else DEFAULT_CATEGORY


def style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = white_bold
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(horizontal="center")


def autosize(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def add_items_sheet(wb: Workbook, items: List[Dict]) -> None:
    ws = wb.active
    ws.title = "Items"
    headers = ["receipt_id", "line_no", "item", "amount_cad", "budget_category", "category_source", "raw_line"]
    ws.append(headers)
    for row in items:
        ws.append([row[h] for h in headers])
    style_header(ws)
    autosize(ws)


def add_receipts_sheet(wb: Workbook, receipts: List[Dict]) -> None:
    ws = wb.create_sheet("Receipts")
    headers = ["receipt_id", "order_id", "receipt_filename", "filename_total_cad", "subtotal_cad", "tax_cad", "total_cad", "parsed_item_count"]
    ws.append(headers)
    for row in receipts:
        ws.append([row[h] for h in headers])
    style_header(ws)
    autosize(ws)


def add_category_summary_sheet(wb: Workbook, item_rows: List[Dict]) -> None:
    ws = wb.create_sheet("Receipt Category Totals")
    categories = list(CATEGORY_ORDER) or [DEFAULT_CATEGORY]
    headers = ["receipt_id"] + categories + ["Total"]
    ws.append(headers)
    receipt_ids = sorted({row["receipt_id"] for row in item_rows})
    for row_idx, receipt_id in enumerate(receipt_ids, start=2):
        ws.cell(row=row_idx, column=1, value=receipt_id)
        for col_idx in range(2, 2 + len(categories)):
            header_ref = ws.cell(row=1, column=col_idx).coordinate
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=f'=SUMIFS(Items!$D:$D,Items!$A:$A,$A{row_idx},Items!$E:$E,{header_ref})',
            )
        ws.cell(
            row=row_idx,
            column=2 + len(categories),
            value=f'=SUM(B{row_idx}:{ws.cell(row=1, column=1 + len(categories)).column_letter}{row_idx})',
        )
    style_header(ws)
    autosize(ws)


def add_notes_sheet(wb: Workbook, notes: List[str]) -> None:
    ws = wb.create_sheet("Notes")
    ws.append(["note"])
    for note in notes:
        ws.append([note])
    style_header(ws)
    autosize(ws)


def _require_valid_schema_name(schema: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", schema or ""):
        raise ValueError(f"Invalid --db-schema value: {schema!r}")
    return schema


def _dsn_contains_password(dsn: str) -> bool:
    if not dsn:
        return False
    key_value_password = re.search(r"(^|\s)password\s*=", dsn, re.IGNORECASE)
    uri_password = re.search(r"://[^/\s:@]+:[^@\s/]*@", dsn)
    return bool(key_value_password or uri_password)


def _resolve_postgres_dsn(raw_dsn: str) -> str:
    dsn = (raw_dsn or "").strip()
    if _dsn_contains_password(dsn):
        raise ValueError(
            "Unsafe --db-dsn: inline password detected. "
            "Use ~/.pgpass (recommended) or PGPASSWORD/HOME_BUDGET_PGPASSWORD env vars."
        )
    return dsn or "dbname=home_budget"


def _connect_postgres(dsn: str):
    dsn = _resolve_postgres_dsn(dsn)
    hb_pg_password = os.environ.get("HOME_BUDGET_PGPASSWORD", "").strip()
    if hb_pg_password and not os.environ.get("PGPASSWORD"):
        os.environ["PGPASSWORD"] = hb_pg_password
    try:
        import psycopg  # type: ignore

        return psycopg.connect(dsn)
    except Exception:
        try:
            import psycopg2  # type: ignore

            return psycopg2.connect(dsn)
        except Exception as e:
            raise RuntimeError(
                "Could not connect to Postgres. Install psycopg (or psycopg2-binary) and verify --db-dsn. "
                "For auth, prefer ~/.pgpass or libpq env vars."
            ) from e


def load_expense_categories_from_postgres(dsn: str, schema: str) -> List[str]:
    schema = _require_valid_schema_name(schema)
    conn = _connect_postgres(dsn)
    sql = f"""
        SELECT category_name
        FROM {schema}.expense_categories
        WHERE is_active
        ORDER BY id
    """
    try:
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()
        return [str(r[0]).strip() for r in rows if r and str(r[0]).strip()]
    finally:
        conn.close()


def date_from_receipt_id(receipt_id: str) -> Optional[datetime.date]:
    m = DATE_PREFIX_RE.match(receipt_id)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%Y%m%d").date()
    except ValueError:
        return None


def total_from_receipt_id(receipt_id: str) -> Optional[float]:
    m = FILENAME_AMOUNT_RE.match(receipt_id)
    if not m:
        return None
    raw = m.group(1).replace("_", ".")
    try:
        return round(float(raw), 2)
    except ValueError:
        return None


def write_sobeys_to_postgres(receipt_rows: List[Dict], item_rows: List[Dict], dsn: str, schema: str) -> Tuple[int, int]:
    schema = _require_valid_schema_name(schema)
    conn = _connect_postgres(dsn)
    written_expenses = 0
    written_items = 0

    items_by_receipt: Dict[str, List[Dict]] = {}
    for item in item_rows:
        items_by_receipt.setdefault(item["receipt_id"], []).append(item)

    expense_sql = f"""
        INSERT INTO {schema}.expenses (
            source, order_id, receipt_filename, order_date, store_name, order_url,
            receipt_item_subtotal, receipt_total_charged, expense_total, raw_payload
        ) VALUES (
            %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s::jsonb
        )
        ON CONFLICT (source, order_id)
        DO UPDATE SET
            receipt_filename = EXCLUDED.receipt_filename,
            order_date = EXCLUDED.order_date,
            store_name = EXCLUDED.store_name,
            order_url = EXCLUDED.order_url,
            receipt_item_subtotal = EXCLUDED.receipt_item_subtotal,
            receipt_total_charged = EXCLUDED.receipt_total_charged,
            expense_total = EXCLUDED.expense_total,
            raw_payload = EXCLUDED.raw_payload
        RETURNING id
    """
    delete_items_sql = f"DELETE FROM {schema}.expense_items WHERE expense_pk = %s"
    item_sql = f"""
        INSERT INTO {schema}.expense_items (
            expense_pk, item_name, budget_category, category_source,
            unit_qty, unit_cost, weight_qty, weight_unit, line_total, original_line_total
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    """

    try:
        conn.autocommit = False
        with conn.cursor() as cur:
            for rec in receipt_rows:
                receipt_id = rec["receipt_id"]
                order_id = str(rec.get("order_id") or receipt_id)
                receipt_items = items_by_receipt.get(receipt_id, [])
                subtotal = rec.get("subtotal_cad")
                total = rec.get("total_cad")
                if total is None and subtotal is not None and rec.get("tax_cad") is not None:
                    total = round(float(subtotal) + float(rec["tax_cad"]), 2)
                payload = {"receipt": rec, "items": receipt_items}

                cur.execute(
                    expense_sql,
                    (
                        "sobeys",
                        order_id,
                        rec.get("receipt_filename"),
                        date_from_receipt_id(receipt_id),
                        "Sobeys",
                        None,
                        subtotal,
                        total,
                        total,
                        json.dumps(payload, ensure_ascii=True),
                    ),
                )
                row = cur.fetchone()
                if row is None:
                    continue
                expense_pk = int(row[0])
                written_expenses += 1

                cur.execute(delete_items_sql, (expense_pk,))

                grouped: Dict[
                    Tuple[str, Optional[str], Optional[str], Optional[float], Optional[float], Optional[str]],
                    Dict[str, Optional[float] | Optional[str]],
                ] = {}
                for item in receipt_items:
                    item_name = normalize_line(str(item.get("item") or ""))
                    unit_cost = item.get("unit_cost")
                    unit_cost_val = float(unit_cost) if unit_cost is not None else None
                    weight_qty_val = float(item["weight_qty"]) if item.get("weight_qty") is not None else None
                    weight_unit_val = str(item.get("weight_unit") or "").lower() or None
                    line_total_val = float(item.get("line_total")) if item.get("line_total") is not None else None
                    key = (
                        item_name.lower(),
                        item.get("budget_category"),
                        item.get("category_source"),
                        unit_cost_val,
                        weight_qty_val,
                        weight_unit_val,
                    )
                    slot = grouped.get(key)
                    if slot is None:
                        grouped[key] = {
                            "item_name": item_name,
                            "budget_category": item.get("budget_category"),
                            "category_source": item.get("category_source"),
                            "unit_qty": float(item["unit_qty"]) if item.get("unit_qty") is not None else None,
                            "unit_cost": unit_cost_val,
                            "weight_qty": weight_qty_val,
                            "weight_unit": weight_unit_val,
                            "line_total": line_total_val,
                        }
                    else:
                        if slot.get("unit_qty") is not None and item.get("unit_qty") is not None:
                            slot["unit_qty"] = float(slot.get("unit_qty") or 0.0) + float(item["unit_qty"])
                        if line_total_val is not None:
                            slot["line_total"] = float(slot.get("line_total") or 0.0) + line_total_val

                for merged in grouped.values():
                    cur.execute(
                        item_sql,
                        (
                            expense_pk,
                            merged["item_name"],
                            merged["budget_category"],
                            merged["category_source"],
                            merged["unit_qty"],
                            merged["unit_cost"],
                            merged["weight_qty"],
                            merged["weight_unit"],
                            merged["line_total"],
                            None,
                        ),
                    )
                    written_items += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return written_expenses, written_items


def main() -> None:
    args = parse_args()
    input_path = Path(args.input_path).expanduser().resolve()
    out_xlsx = Path(args.out_xlsx).expanduser().resolve()

    try:
        db_categories = load_expense_categories_from_postgres(args.db_dsn, args.db_schema)
        if db_categories:
            CATEGORY_ORDER[:] = db_categories
            print(f"Loaded categories from {args.db_schema}.expense_categories: {len(db_categories)} active", flush=True)
        else:
            print("Warning: no active rows in expense_categories; using built-in categories.", flush=True)
    except Exception as e:
        print(f"Warning: could not load categories from DB; using built-in categories. {type(e).__name__}: {e}", flush=True)

    if not input_path.exists():
        raise FileNotFoundError(f"Missing input path: {input_path}")

    AI_ENGINE.load_suggestions()
    started = time.perf_counter()
    files = extract_inputs(input_path)
    if not files:
        raise RuntimeError(f"No receipt files found under: {input_path}")

    all_items: List[Dict] = []
    receipt_rows: List[Dict] = []
    notes: List[str] = []

    for file_path in files:
        receipt_id = file_path.stem
        text = document_text(file_path)
        alt_text = document_text_alt(file_path)
        items = parse_items(receipt_id, text, alt_text=alt_text)
        subtotal, tax, total = parse_receipt_amounts(text)
        filename_total = total_from_receipt_id(receipt_id)
        if total is None and filename_total is not None:
            total = filename_total
        if not items and total is not None:
            items = [
                {
                    "receipt_id": receipt_id,
                    "line_no": 1,
                    "item": "UNPARSED SOBEYS RECEIPT",
                    "amount_cad": float(total),
                    "budget_category": fallback_category(),
                    "category_source": "ocr_fallback",
                    "raw_line": "OCR fallback: no item lines detected; using receipt total.",
                }
            ]
        elif not items:
            fallback_total = filename_total if filename_total is not None else 0.0
            items = [
                {
                    "receipt_id": receipt_id,
                    "line_no": 1,
                    "item": "UNPARSED SOBEYS RECEIPT",
                    "amount_cad": fallback_total,
                    "budget_category": fallback_category(),
                    "category_source": "ocr_no_items_filename_total" if filename_total is not None else "ocr_no_items",
                    "raw_line": (
                        f"OCR fallback: no item lines; using filename total {filename_total:.2f}."
                        if filename_total is not None
                        else "OCR fallback: no item lines or total detected."
                    ),
                }
            ]
        elif total is not None:
            parsed_sum = round(sum(float(i.get("amount_cad") or 0.0) for i in items), 2)
            mismatch = round(float(total) - parsed_sum, 2)
            if abs(mismatch) > max(2.0, float(total) * 0.20):
                items.append(
                    {
                        "receipt_id": receipt_id,
                        "line_no": 999999,
                        "item": "UNPARSED SOBEYS RECEIPT REMAINDER",
                        "amount_cad": mismatch,
                        "budget_category": fallback_category(),
                        "category_source": "ocr_remainder",
                        "raw_line": (
                            "OCR remainder fallback: parsed item sum does not reconcile to receipt total. "
                            f"parsed_sum={parsed_sum:.2f}, total={float(total):.2f}"
                        ),
                    }
                )
        all_items.extend(items)

        receipt_rows.append(
            {
                "receipt_id": receipt_id,
                "order_id": receipt_id,
                "receipt_filename": file_path.name,
                "subtotal_cad": subtotal,
                "tax_cad": tax,
                "total_cad": total,
                "filename_total_cad": filename_total,
                "parsed_item_count": len(items),
            }
        )

    all_items.sort(key=lambda x: (x["receipt_id"], x["line_no"]))
    receipt_rows.sort(key=lambda x: x["receipt_id"])

    wb = Workbook()
    add_items_sheet(wb, all_items)
    add_receipts_sheet(wb, receipt_rows)
    add_category_summary_sheet(wb, all_items)
    notes.append(f"Input path: {input_path}")
    notes.append(f"Receipts parsed: {len(receipt_rows)}")
    notes.append(f"Items parsed: {len(all_items)}")
    notes.append(
        "AI Status: "
        f"enabled={AI_ENGINE.enabled}, "
        f"disabled_runtime={AI_ENGINE.runtime_disabled}, "
        f"calls={AI_ENGINE.stats['calls']}, "
        f"batch_items={AI_ENGINE.stats['batch_items']}, "
        f"ai_cache_hits={AI_ENGINE.stats['ai_cache_hits']}, "
        f"success={AI_ENGINE.stats['success']}, "
        f"api_errors={AI_ENGINE.stats['api_errors']}."
    )
    add_notes_sheet(wb, notes)
    wb.save(out_xlsx)
    AI_ENGINE.save_suggestions()

    if args.write_db:
        written_expenses, written_items = write_sobeys_to_postgres(
            receipt_rows=receipt_rows,
            item_rows=all_items,
            dsn=args.db_dsn,
            schema=args.db_schema,
        )
        print(
            f"Wrote Postgres: schema={args.db_schema}, expenses={written_expenses}, expense_items={written_items}",
            flush=True,
        )

    print(f"Wrote workbook: {out_xlsx}", flush=True)
    print(f"Runtime timing: total={time.perf_counter() - started:.2f}s", flush=True)


if __name__ == "__main__":
    main()
