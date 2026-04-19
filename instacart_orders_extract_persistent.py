#!/usr/bin/env python3
"""
Instacart order extractor using Playwright persistent profile.

Why persistent profile:
- Works with 2FA
- Works with Google login
- No cookie export/import needed
- Stable between runs

Usage:
  python instacart_orders_extract_persistent.py --headful --interactive-login

First run:
- A browser window opens with a persistent profile directory.
- Complete login/2FA manually if needed.
- Script continues and scrapes order history.

Outputs:
- instacart_orders.json
- instacart_orders_items.csv
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import time
from datetime import datetime
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from openpyxl import Workbook
from playwright.sync_api import BrowserContext, Page, TimeoutError as PlaywrightTimeoutError, sync_playwright
from budget_category_ai import AICategoryEngine
from budget_category_logic import CATEGORY_ORDER, DEFAULT_CATEGORY, canonicalize_category, deterministic_category, normalize_for_match

INSTACART_BASE = "https://www.instacart.ca"
DEFAULT_ORDERS_URLS = [
    "https://www.instacart.ca/store/account/family_orders",
    "https://www.instacart.ca/store/account/orders",
    "https://www.instacart.ca/orders",
    "https://www.instacart.com/store/account/family_orders",
    "https://www.instacart.com/store/account/orders",
    "https://www.instacart.com/orders",
]
AI_SUGGESTIONS_PATH = Path('/Users/paul/Documents/Finance/Budget/2026/Groceries/ai_category_cache.json')
DEBUG_ITEM_PATTERN = os.environ.get("INSTACART_DEBUG_ITEM", "").strip().lower()


def debug_item_match(*texts: Optional[str]) -> bool:
    if not DEBUG_ITEM_PATTERN:
        return False
    for t in texts:
        if t and DEBUG_ITEM_PATTERN in t.lower():
            return True
    return False


def debug_item_log(message: str) -> None:
    if DEBUG_ITEM_PATTERN:
        print(f"[DEBUG item:{DEBUG_ITEM_PATTERN}] {message}", flush=True)


@dataclass
class OrderItem:
    name: str
    quantity: Optional[str]
    unit_price: Optional[str]
    total_price: Optional[str]
    weight_qty: Optional[float] = None
    weight_unit: Optional[str] = None
    original_total_price: Optional[str] = None
    budget_category: Optional[str] = None
    category_source: Optional[str] = None
    quantity_locked: bool = False


@dataclass
class OrderRecord:
    order_url: str
    order_id: Optional[str]
    order_date: Optional[str]
    store_name: Optional[str]
    order_total: Optional[str]
    receipt_item_subtotal: Optional[str]
    receipt_tip: Optional[str]
    receipt_service_fee: Optional[str]
    receipt_recycling_fee: Optional[str]
    receipt_service_fee_tax: Optional[str]
    receipt_gst: Optional[str]
    receipt_pst: Optional[str]
    receipt_discount_total: Optional[str]
    receipt_total_charged: Optional[str]
    item_count: Optional[int]
    items: List[OrderItem]
    raw_page_title: Optional[str]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract Instacart orders with Playwright persistent profile.")
    parser.add_argument(
        "--profile-dir",
        default=str(Path.home() / ".config" / "instacart-playwright-profile"),
        help="Directory for persistent browser profile.",
    )
    parser.add_argument(
        "--out-json",
        default="instacart_orders.json",
        help="Output JSON path.",
    )
    parser.add_argument(
        "--out-csv",
        default="instacart_orders_items.csv",
        help="Output CSV path (flattened items).",
    )
    parser.add_argument(
        "--out-xlsx",
        default="instacart_orders_items.xlsx",
        help="Output Excel path (cleaned flattened items + orders summary).",
    )
    parser.add_argument(
        "--max-orders",
        type=int,
        default=0,
        help="Maximum orders to scrape (0 means all discovered).",
    )
    parser.add_argument(
        "--headful",
        action="store_true",
        help="Run browser in headful mode (recommended for login).",
    )
    parser.add_argument(
        "--interactive-login",
        action="store_true",
        help="Wait for manual login if the session is not already authenticated.",
    )
    parser.add_argument(
        "--timeout-ms",
        type=int,
        default=25000,
        help="Playwright timeout in milliseconds.",
    )
    parser.add_argument(
        "--scroll-rounds",
        type=int,
        default=40,
        help="How many load-more scroll passes to run on orders page.",
    )
    parser.add_argument(
        "--month",
        default="",
        help="Filter output to one month in YYYY-MM format (e.g., 2026-04).",
    )
    parser.add_argument(
        "--write-db",
        action="store_true",
        help="Write parsed canonical expense records into Postgres.",
    )
    parser.add_argument(
        "--db-dsn",
        default=os.environ.get("HOME_BUDGET_PG_DSN", ""),
        help=(
            "Postgres DSN for --write-db. Do not include password here. "
            "Use ~/.pgpass or PGPASSWORD/HOME_BUDGET_PGPASSWORD env vars. "
            "Default: env HOME_BUDGET_PG_DSN or dbname=home_budget."
        ),
    )
    parser.add_argument(
        "--db-schema",
        default="budget",
        help="Target schema for --write-db (default: budget).",
    )
    parser.add_argument(
        "--default-store-name",
        default="",
        help="Fallback store name to use when a parsed order has no store name (example: Costco).",
    )
    return parser.parse_args()


def launch_context(profile_dir: Path, headful: bool, timeout_ms: int) -> Tuple[Any, BrowserContext]:
    pw = sync_playwright().start()
    context = pw.chromium.launch_persistent_context(
        user_data_dir=str(profile_dir),
        headless=not headful,
        viewport={"width": 1440, "height": 980},
        args=["--disable-blink-features=AutomationControlled"],
    )
    context.set_default_timeout(timeout_ms)
    context.set_default_navigation_timeout(timeout_ms)
    return pw, context


def get_or_create_page(context: BrowserContext) -> Page:
    if context.pages:
        return context.pages[0]
    return context.new_page()


def looks_logged_in(page: Page) -> bool:
    url = page.url
    if "/login" in url or "/accounts/login" in url:
        return False
    if "/orders" in url or "/account" in url:
        return True

    try:
        text = page.locator("body").inner_text(timeout=2500).lower()
    except Exception:
        return False

    logged_in_markers = [
        "your orders",
        "order history",
        "account settings",
        "log out",
        "logout",
    ]
    if any(marker in text for marker in logged_in_markers):
        return True

    login_markers = [
        "sign in",
        "continue with google",
        "enter your phone number",
    ]
    if any(marker in text for marker in login_markers):
        if "your orders" not in text and "order history" not in text:
            return False
    return True


def wait_for_manual_login(page: Page, timeout_seconds: int = 600) -> None:
    print(
        "Complete login/2FA in browser, then press Enter here to continue immediately.",
        flush=True,
    )
    try:
        input()
        return
    except EOFError:
        pass

    started = time.time()
    while time.time() - started < timeout_seconds:
        if looks_logged_in(page):
            return
        print("Waiting for login/2FA completion in browser...", flush=True)
        time.sleep(3)
    raise TimeoutError("Login was not completed within timeout window.")


def open_orders_page(page: Page) -> None:
    for url in DEFAULT_ORDERS_URLS:
        try:
            page.goto(url, wait_until="domcontentloaded")
            time.sleep(1)
            if "orders" in page.url:
                return
        except PlaywrightTimeoutError:
            continue

    page.goto(INSTACART_BASE, wait_until="domcontentloaded")


def incremental_scroll(page: Page, rounds: int) -> None:
    last_link_count = -1
    stable_rounds = 0
    max_rounds = max(1, rounds)

    for _ in range(max_rounds):
        # Trigger lazy loading.
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        time.sleep(0.9)

        # Some account pages use explicit pagination/load buttons instead of pure infinite scroll.
        click_labels = (
            "show more",
            "load more",
            "see more",
            "view more",
            "older orders",
            "next",
        )
        for label in click_labels:
            try:
                btn = page.get_by_text(label, exact=False).first
                if btn.count() and btn.is_visible():
                    btn.click(timeout=1200)
                    time.sleep(0.5)
            except Exception:
                pass

        # Small upward nudge helps virtualized lists render next chunk.
        try:
            page.mouse.wheel(0, -500)
            page.mouse.wheel(0, 1200)
        except Exception:
            pass

        # Early stop if no new order links appear after several rounds.
        try:
            link_count = len(extract_order_links(page))
        except Exception:
            link_count = -1
        if link_count <= last_link_count:
            stable_rounds += 1
        else:
            stable_rounds = 0
            last_link_count = link_count
        if stable_rounds >= 6:
            break


def normalize_order_url(url: str) -> Optional[str]:
    if not url:
        return None
    if url.startswith("/"):
        url = INSTACART_BASE + url
    if not url.startswith("http"):
        return None

    # Common Instacart order paths include both plural and singular variants.
    if "/orders/" not in url and "/order/" not in url:
        return None

    # Strip fragments/query for stable dedupe
    return url.split("#", 1)[0].split("?", 1)[0]


def extract_order_links(page: Page) -> List[str]:
    links: List[str] = page.evaluate(
        """
        () => {
          const out = [];
          for (const a of Array.from(document.querySelectorAll('a[href]'))) {
            out.push(a.href || a.getAttribute('href') || '');
          }
          return out;
        }
        """
    )

    uniq: Set[str] = set()
    ordered: List[str] = []
    for raw in links:
        norm = normalize_order_url(raw)
        if norm and norm not in uniq:
            uniq.add(norm)
            ordered.append(norm)
    return ordered


def parse_store_hint_text(block_text: str) -> Optional[str]:
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in block_text.splitlines() if ln.strip()]
    if not lines:
        return None

    stop_words = (
        "your order was delivered",
        "delivered",
        "receipt",
        "reorder",
        "order total",
        "items subtotal",
        "delivery details",
        "family order",
        "instacart",
        "$",
    )
    for ln in lines:
        ll = ln.lower()
        if any(tok in ll for tok in stop_words):
            continue
        if re.search(r"\d", ln):
            continue
        if len(ln) < 2 or len(ln) > 60:
            continue
        return ln
    return None


def extract_order_store_hints(page: Page) -> Dict[str, str]:
    raw_blocks: List[Dict[str, str]] = page.evaluate(
        """
        () => {
          const out = [];
          for (const a of Array.from(document.querySelectorAll('a[href*="/orders/"]'))) {
            const href = a.href || a.getAttribute('href') || '';
            if (!href) continue;

            let blockText = '';
            let n = a;
            for (let depth = 0; depth < 8 && n; depth += 1, n = n.parentElement) {
              const txt = (n.innerText || '').trim();
              if (!txt) continue;
              if (/your order was delivered/i.test(txt) || /delivery details/i.test(txt)) {
                blockText = txt;
                break;
              }
              if (!blockText && txt.length > 10 && txt.length < 1200) {
                blockText = txt;
              }
            }
            out.push({ href, text: blockText || (a.innerText || '') });
          }
          return out.slice(0, 1200);
        }
        """
    )

    hints: Dict[str, str] = {}
    for block in raw_blocks:
        norm = normalize_order_url(block.get("href", ""))
        if not norm:
            continue
        hint = parse_store_hint_text(block.get("text", ""))
        hint = clean_store_name(hint) if hint else None
        if hint and norm not in hints:
            hints[norm] = hint
    return hints


def wait_for_meaningful_body_text(page: Page, timeout_seconds: float = 10.0) -> str:
    deadline = time.time() + max(0.5, timeout_seconds)
    best = ""
    while time.time() < deadline:
        try:
            text = page.locator("body").inner_text(timeout=2000).strip()
        except Exception:
            text = ""
        if len(text) > len(best):
            best = text
        ll = text.lower()
        if text and (
            len(text) > 200
            or any(tok in ll for tok in ("items subtotal", "order total", "receipt", "delivered", "your order"))
        ):
            return text
        time.sleep(0.4)
    return best


def capture_receipt_text(page: Page, timeout_seconds: float = 8.0) -> str:
    deadline = time.time() + max(1.0, timeout_seconds)
    best = ""

    # Expand "Receipt" accordion when present.
    try:
        receipt_btn = page.get_by_text("Receipt", exact=False).first
        if receipt_btn.count() and receipt_btn.is_visible():
            receipt_btn.click(timeout=1200)
            time.sleep(0.35)
    except Exception:
        pass

    while time.time() < deadline:
        try:
            page.mouse.wheel(0, 1800)
        except Exception:
            pass
        time.sleep(0.35)

        try:
            txt = page.locator("body").inner_text(timeout=1800).strip()
        except Exception:
            txt = ""
        if len(txt) > len(best):
            best = txt

        ll = txt.lower()
        if (
            "item subtotal" in ll
            and "service fee" in ll
            and ("total" in ll or "total charged" in ll)
        ):
            break

    return best


def candidate_order_detail_urls(page: Page, order_url: str) -> List[str]:
    out: List[str] = []
    seen: Set[str] = set()
    order_id = (order_url.rstrip("/").split("/")[-1] or "").strip()

    suffix_candidates = ("receipt", "receipts", "details", "order_receipt", "order-receipt")
    for suffix in suffix_candidates:
        candidate = f"{order_url.rstrip('/')}/{suffix}"
        if candidate not in seen:
            seen.add(candidate)
            out.append(candidate)

    if "/store/orders/" in order_url:
        alt = order_url.replace("/store/orders/", "/store/order_receipts/")
        if alt not in seen:
            seen.add(alt)
            out.append(alt)

    try:
        hrefs: List[str] = page.evaluate(
            """
            () => Array.from(document.querySelectorAll('a[href]'))
              .map(a => a.href || a.getAttribute('href') || '')
            """
        )
    except Exception:
        hrefs = []

    ranked: List[Tuple[int, str]] = []
    for raw in hrefs:
        norm = normalize_order_url(raw)
        if not norm:
            continue
        ll = norm.lower()
        if order_id and order_id not in ll:
            continue
        score = 0
        if "receipt" in ll:
            score += 6
        if "detail" in ll:
            score += 3
        if "/orders/" in ll:
            score += 1
        if score <= 0:
            continue
        ranked.append((score, norm))

    for _, url in sorted(ranked, key=lambda x: (-x[0], x[1])):
        if url not in seen:
            seen.add(url)
            out.append(url)
    return out


def text_or_none(page: Page, selector: str) -> Optional[str]:
    try:
        loc = page.locator(selector).first
        if loc.count() == 0:
            return None
        text = loc.inner_text(timeout=2000).strip()
        return text if text else None
    except Exception:
        return None


def parse_order_id(title: str, body_text: str, order_url: str) -> Optional[str]:
    for source in (title, body_text):
        m_num = re.search(r"\bOrder\s*#\s*([0-9]{8,})\b", source, re.I)
        if m_num:
            return m_num.group(1)

    # Fallback to URL token only if no numeric order id is visible on receipt.
    m = re.search(r"/orders/([A-Za-z0-9\-]+)", order_url)
    if m:
        return m.group(1)
    return None


def parse_order_date(body_text: str) -> Optional[str]:
    patterns = [
        r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\s+\d{1,2},\s+\d{4}\b",
        r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\s+\d{1,2}\b",
        r"\b\d{4}-\d{2}-\d{2}\b",
        r"\b\d{1,2}/\d{1,2}/\d{4}\b",
    ]
    for pat in patterns:
        m = re.search(pat, body_text, re.I)
        if m:
            return m.group(0)
    return None


def parse_order_date_from_order_id(order_id: Optional[str]) -> Optional[str]:
    if not order_id:
        return None
    m = re.match(r"^(\d{2})(\d{2})\d{6,}$", order_id)
    if not m:
        return None

    month = int(m.group(1))
    day = int(m.group(2))
    if month < 1 or month > 12:
        return None

    now = datetime.now()
    year = now.year - 1 if month > now.month else now.year
    try:
        dt = datetime(year, month, day)
    except ValueError:
        return None
    return dt.strftime("%Y-%m-%d")


def month_matches(order_date: Optional[str], target_month: str) -> bool:
    if not target_month:
        return True
    if not order_date:
        # Keep undated orders when month filtering; date parsing can fail on some pages.
        return True

    m = re.fullmatch(r"(\d{4})-(\d{2})", target_month)
    if not m:
        raise ValueError(f"Invalid --month format: {target_month} (expected YYYY-MM)")
    target_year = int(m.group(1))
    target_mon = int(m.group(2))
    if target_mon < 1 or target_mon > 12:
        raise ValueError(f"Invalid --month value: {target_month} (month must be 01-12)")

    text = order_date.strip()
    parse_formats = [
        "%b %d, %Y",
        "%B %d, %Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ]
    for fmt in parse_formats:
        try:
            dt = datetime.strptime(text, fmt)
            return dt.year == target_year and dt.month == target_mon
        except ValueError:
            continue

    # Accept month names without year (e.g., "Apr 18") by month only.
    short_map = {
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
        "jul": 7, "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12,
    }
    mm = re.match(r"([A-Za-z]+)\s+\d{1,2}$", text)
    if mm:
        token = mm.group(1).lower().rstrip(".")
        mon = None
        for key, val in short_map.items():
            if token.startswith(key):
                mon = val
                break
        return mon == target_mon

    return False


def parse_order_total(body_text: str) -> Optional[str]:
    lines = [ln.strip() for ln in body_text.splitlines() if ln.strip()]
    if not lines:
        return None

    total_label_re = re.compile(r"\b(order\s+total|total\s+charged|amount\s+paid|charged)\b", re.I)
    excluded_line_re = re.compile(
        r"\b(for you|coupon|discount|saved|reward|promo|promotion|credit|recycling fee|tip)\b",
        re.I,
    )
    money_re = re.compile(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?")

    # First pass: explicit total labels on the same line.
    for ln in lines:
        if not total_label_re.search(ln):
            continue
        if excluded_line_re.search(ln):
            continue
        money = money_re.findall(ln)
        if money:
            return money[-1].replace(" ", "")

    # Second pass: label followed by amount on nearby next line.
    for idx, ln in enumerate(lines):
        if not total_label_re.search(ln):
            continue
        if excluded_line_re.search(ln):
            continue
        if idx + 1 < len(lines):
            nxt = lines[idx + 1]
            m = money_re.search(nxt)
            if m:
                return m.group(0).replace(" ", "")

    # Fallback: choose the largest money value from non-promotional lines.
    candidates: List[str] = []
    for ln in lines:
        if excluded_line_re.search(ln):
            continue
        candidates.extend(money_re.findall(ln))

    if not candidates:
        candidates = money_re.findall(body_text)
    if not candidates:
        return None

    normalized = [m.replace(" ", "").replace("$", "") for m in candidates]
    try:
        top = max(normalized, key=lambda x: float(x.replace(",", "")))
        return f"${top}"
    except Exception:
        return candidates[0].replace(" ", "")


def parse_receipt_breakdown(body_text: str) -> Dict[str, Optional[str]]:
    raw_lines = [re.sub(r"\s+", " ", ln).strip() for ln in body_text.splitlines() if ln.strip()]
    lines: List[str] = []
    seen_lines: Set[str] = set()
    for ln in raw_lines:
        if ln in seen_lines:
            continue
        seen_lines.add(ln)
        lines.append(ln)
    if not lines:
        return {
            "item_subtotal": None,
            "tip": None,
            "service_fee": None,
            "recycling_fee": None,
            "service_fee_tax": None,
            "gst": None,
            "pst": None,
            "discount_total": None,
            "total_charged": None,
        }

    money_re = re.compile(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?")
    signed_money_re = re.compile(r"(?P<prefix>-)?\$\s?(?P<num>\d+(?:,\d{3})*(?:\.\d{2})?)(?P<suffix>-)?")
    discount_re = re.compile(r"\b(discount|coupon|promo|promotion|credit|adjustment)\b", re.I)

    def parse_signed_money(text: str, label_hint: str = "") -> Optional[float]:
        m = signed_money_re.search(text)
        if not m:
            return None
        try:
            value = float(m.group("num").replace(",", ""))
        except Exception:
            return None
        is_negative = bool(m.group("prefix") or m.group("suffix"))
        if not is_negative and discount_re.search(label_hint):
            is_negative = True
        return -value if is_negative else value

    def format_signed_money(value: float) -> str:
        return f"-${abs(value):.2f}" if value < 0 else f"${value:.2f}"

    patterns: Dict[str, re.Pattern[str]] = {
        "item_subtotal": re.compile(r"\b(item|items)\s+subtotal\b", re.I),
        "tip": re.compile(r"^tip\b", re.I),
        "service_fee": re.compile(r"^service fee\b(?!\s*tax)", re.I),
        "recycling_fee": re.compile(r"^recycling fee\b", re.I),
        "service_fee_tax": re.compile(r"^service fee tax\b", re.I),
        "gst": re.compile(r"\bgoods and services tax\b|\bgst\b", re.I),
        "pst": re.compile(r"\bprovincial sales tax\b|\bpst\b", re.I),
        "total_charged": re.compile(r"^total\b|order total|total charged|amount paid|charged\b", re.I),
    }

    out: Dict[str, Optional[str]] = {k: None for k in patterns.keys()}
    for idx, line in enumerate(lines):
        for key, pat in patterns.items():
            if out[key] is not None:
                continue
            if not pat.search(line):
                continue
            vals = money_re.findall(line)
            if not vals and idx + 1 < len(lines):
                vals = money_re.findall(lines[idx + 1])
            if vals:
                out[key] = vals[-1].replace(" ", "")
                continue
            if re.search(r"\bfree\b", line, re.I):
                out[key] = "$0.00"
                continue
            if idx + 1 < len(lines) and re.search(r"\bfree\b", lines[idx + 1], re.I):
                out[key] = "$0.00"

    # Fallback: when labels and amounts are split across separate short lines in receipt grid.
    for idx, line in enumerate(lines):
        ll = line.lower()
        key: Optional[str] = None
        if re.search(r"\b(item|items)\s+subtotal\b", ll):
            key = "item_subtotal"
        elif re.fullmatch(r"tip", ll):
            key = "tip"
        elif re.fullmatch(r"service fee", ll):
            key = "service_fee"
        elif re.fullmatch(r"recycling fee", ll):
            key = "recycling_fee"
        elif re.fullmatch(r"service fee tax", ll):
            key = "service_fee_tax"
        elif re.search(r"goods and services tax|\bgst\b", ll):
            key = "gst"
        elif re.search(r"provincial sales tax|\bpst\b", ll):
            key = "pst"
        elif re.fullmatch(r"total", ll):
            key = "total_charged"
        if not key or out.get(key):
            continue

        for j in range(idx + 1, min(len(lines), idx + 5)):
            vals = money_re.findall(lines[j])
            if vals:
                out[key] = vals[-1].replace(" ", "")
                break

    # Discounts should come from receipt rows, not page-level "you saved" summaries.
    receipt_start = 0
    receipt_end = len(lines)
    for idx, line in enumerate(lines):
        if re.search(r"\b(item|items)\s+subtotal\b", line, re.I):
            receipt_start = idx
            break
    for idx in range(receipt_start, len(lines)):
        if re.search(r"^total\b|order total|total charged|amount paid|charged\b", lines[idx], re.I):
            receipt_end = idx + 1
            break
    receipt_lines = lines[receipt_start:receipt_end]

    discount_total = 0.0
    found_discount = False
    used_amount_keys: Set[str] = set()
    for idx, line in enumerate(receipt_lines):
        if not discount_re.search(line):
            continue
        if re.search(r"\byou saved\b|\bsavings?\b", line, re.I):
            continue
        amount: Optional[float] = None
        amount_key: Optional[str] = None
        for j in range(idx, min(len(receipt_lines), idx + 3)):
            amount = parse_signed_money(receipt_lines[j], label_hint=line)
            if amount is not None:
                amount_key = re.sub(r"\s+", "", receipt_lines[j].lower())
                break
        if amount is None:
            continue
        if amount_key and amount_key in used_amount_keys:
            continue
        if amount_key:
            used_amount_keys.add(amount_key)
        if amount > 0:
            amount = -amount
        discount_total += amount
        found_discount = True
    if found_discount:
        out["discount_total"] = format_signed_money(discount_total)
    return out


def parse_store_name(body_text: str) -> Optional[str]:
    lines = [ln.strip() for ln in body_text.splitlines() if ln.strip()]

    for ln in lines:
        m = re.search(r"\bfrom\s+(.+)$", ln, re.I)
        if m:
            candidate = m.group(1).strip(" .")
            if 2 <= len(candidate) <= 80:
                return candidate

    banned = (
        "instacart",
        "order",
        "receipt",
        "delivered",
        "delivering",
        "substitution",
        "item",
        "total",
        "tip",
        "tax",
        "$",
    )
    for ln in lines:
        ll = ln.lower()
        if any(tok in ll for tok in banned):
            continue
        if len(ln) < 3 or len(ln) > 80:
            continue
        if re.search(r"\d", ln):
            continue
        return ln
    return None


def parse_store_name_from_dom(page: Page) -> Optional[str]:
    try:
        candidate = page.evaluate(
            """
            () => {
              const noise = new Set([
                'instacart', 'all stores', 'your order was delivered', 'delivery details',
                'family order', 'receipt', 'your items', 'found', 'skip navigation'
              ]);

              const pickLine = (lines) => {
                for (const raw of lines) {
                  const line = (raw || '').trim();
                  if (!line) continue;
                  const ll = line.toLowerCase();
                  if (noise.has(ll)) continue;
                  if (/^found\\s*\\(\\d+\\)$/i.test(line)) continue;
                  if (/\\d/.test(line)) continue;
                  if (line.length < 2 || line.length > 60) continue;
                  return line;
                }
                return null;
              };

              // Best signal: order status card usually contains store line above "Your order was delivered".
              for (const n of Array.from(document.querySelectorAll('section, article, div, aside'))) {
                const txt = (n.innerText || '').trim();
                if (!txt) continue;
                if (!/your order was delivered/i.test(txt)) continue;
                const lines = txt.split(/\\r?\\n+/);
                const picked = pickLine(lines);
                if (picked) return picked;
              }

              // Fallback selectors used by many retailer/order templates.
              const selectors = [
                '[data-testid*="store"]',
                '[data-testid*="retailer"]',
                '[data-testid*="merchant"]',
                '[aria-label*="store"]',
              ];
              for (const sel of selectors) {
                for (const n of Array.from(document.querySelectorAll(sel))) {
                  const txt = (n.innerText || n.getAttribute('aria-label') || '').trim();
                  if (!txt) continue;
                  const picked = pickLine(txt.split(/\\r?\\n+/));
                  if (picked) return picked;
                }
              }

              // Search box sometimes exposes retailer as "Search Costco".
              for (const el of Array.from(document.querySelectorAll('input[placeholder], input[aria-label], [role="searchbox"]'))) {
                const txt = (
                  el.getAttribute('placeholder')
                  || el.getAttribute('aria-label')
                  || el.getAttribute('value')
                  || ''
                ).trim();
                if (!txt) continue;
                const m = txt.match(/^search\\s+(.+)$/i);
                if (m && m[1]) return m[1].trim();
              }

              // Retailer links often include /store/<slug>; prefer human-readable label on the same node.
              for (const a of Array.from(document.querySelectorAll('a[href*="/store/"]'))) {
                const href = a.getAttribute('href') || '';
                const t = (a.innerText || a.getAttribute('aria-label') || a.getAttribute('title') || '').trim();
                if (t && !/^all stores$/i.test(t) && !/^instacart$/i.test(t)) return t;
                const m = href.match(/\\/store\\/([a-z0-9-]+)/i);
                if (m && m[1] && !/^orders?$/i.test(m[1])) {
                  return m[1].replace(/-/g, ' ');
                }
              }

              // Store logos can carry alt/title text (e.g., "Costco").
              for (const n of Array.from(document.querySelectorAll('img[alt], img[title]'))) {
                const txt = (n.getAttribute('alt') || n.getAttribute('title') || '').trim();
                if (!txt) continue;
                const ll = txt.toLowerCase();
                if (noise.has(ll)) continue;
                if (/logo|image|photo/.test(ll)) continue;
                if (txt.length >= 2 && txt.length <= 40 && !/\\d/.test(txt)) return txt;
              }
              return null;
            }
            """
        )
    except Exception:
        return None

    if not candidate:
        return None
    value = re.sub(r"\s+", " ", str(candidate)).strip()
    return value if value else None


def wait_for_store_name(page: Page, timeout_seconds: float = 12.0) -> Optional[str]:
    deadline = time.time() + max(1.0, timeout_seconds)
    while time.time() < deadline:
        candidate = (
            parse_store_name_from_dom(page)
            or text_or_none(page, "[data-testid*='store']")
            or text_or_none(page, "[data-testid*='retailer']")
            or text_or_none(page, "[data-testid*='merchant']")
            or text_or_none(page, "h1")
            or text_or_none(page, "h2")
        )
        cleaned = clean_store_name(candidate)
        if cleaned:
            return cleaned
        time.sleep(0.35)
    return None


def clean_store_name(store_name: Optional[str]) -> Optional[str]:
    if not store_name:
        return None
    value = re.sub(r"\s+", " ", store_name).strip()
    if not value:
        return None

    ll = value.lower()
    if "delivered your order" in ll:
        return None
    if "instacart" in ll:
        return None
    m_search = re.match(r"^search\s+(.+)$", value, re.I)
    if m_search:
        value = m_search.group(1).strip()
        ll = value.lower()
        if not value:
            return None
    ui_noise_patterns = [
        r"^skip navigation$",
        r"^all stores$",
        r"^receipt$",
        r"^delivery details$",
        r"^family order$",
        r"^your items$",
        r"^search\b",
        r"^pickup unavailable$",
    ]
    if any(re.search(pat, ll) for pat in ui_noise_patterns):
        return None

    # Remove likely section headers misdetected as stores.
    if value == value.upper() and len(value.split()) <= 4:
        section_words = {
            "canned",
            "goods",
            "produce",
            "dairy",
            "bakery",
            "frozen",
            "snacks",
            "beverages",
            "meat",
            "seafood",
            "deli",
            "pantry",
        }
        tokens = {t.lower() for t in re.findall(r"[A-Za-z]+", value)}
        if tokens and tokens.issubset(section_words):
            return None
    return value


def infer_store_name_from_items(items: List[OrderItem]) -> Optional[str]:
    if not items:
        return None
    names = " ".join((item.name or "").lower() for item in items)
    # Instacart Costco orders frequently include Kirkland-branded SKUs.
    if "kirkland" in names:
        return "Costco"
    return None


def parse_money_value(value: Optional[str]) -> Optional[float]:
    if not value:
        return None
    m = re.search(r"(?P<prefix>-)?\$?\s*(?P<num>[0-9]+(?:,[0-9]{3})*(?:\.[0-9]{2})?)(?P<suffix>-)?", value)
    if not m:
        return None
    try:
        parsed = float(m.group("num").replace(",", ""))
        if m.group("prefix") or m.group("suffix"):
            parsed = -parsed
        return parsed
    except Exception:
        return None


def format_money(value: Optional[float]) -> Optional[str]:
    if value is None:
        return None
    return f"${value:.2f}"


def derive_unit_qty(item: OrderItem) -> Optional[float]:
    if item.weight_unit:
        return None
    if item.quantity:
        m = re.search(r"\d+(?:\.\d+)?", str(item.quantity))
        if m:
            try:
                return float(m.group(0))
            except ValueError:
                pass

    unit = parse_money_value(item.unit_price)
    total = parse_money_value(item.total_price)
    if unit is None or total is None or unit <= 0:
        return None
    qty = total / unit
    if qty <= 0:
        return None
    rounded_int = int(round(qty))
    if abs(qty - rounded_int) < 0.02:
        return float(rounded_int)
    return round(qty, 3)


def derive_weight_qty(item: OrderItem) -> Optional[float]:
    if item.weight_qty is not None and item.weight_qty > 0:
        return round(item.weight_qty, 3)
    if not item.weight_unit:
        return None
    unit = parse_money_value(item.unit_price)
    total = parse_money_value(item.total_price)
    if unit is None or total is None or unit <= 0:
        return None
    return round(total / unit, 3)


def format_unit_qty_for_csv(value: Optional[float]) -> Optional[str]:
    if value is None:
        return None
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return f"{value:.3f}".rstrip("0").rstrip(".")


def format_weight_unit_for_csv(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    return value.lower()


def normalize_quantity(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().lower()
    if re.search(r"\b\d+(?:\.\d+)?\s*(kg|g|lb|lbs|oz)\b", text):
        return None
    m = re.search(r"(\d+)", text)
    if not m:
        return None
    return str(int(m.group(1)))


def is_summary_item_name(name: str) -> bool:
    n = name.strip().lower()
    if not n:
        return True
    summary_patterns = [
        r"^receipt$",
        r"^for you,?$",
        r"^for you,\s*for a friend$",
        r"^,?\s*for a friend$",
        r"friends can get .*terms apply\.?$",
        r"\bterms apply\.?$",
        r"^recycling fee$",
        r"^\$?\d+(?:\.\d{2})?\s*earned\.?$",
        r"\bearned\.?$",
        r"\breward(?:s)?\b",
        r"^\$?\d+(?:\.\d{2})?\s+for you\b",
        r"\bfor you,\s*\$?\d+(?:\.\d{2})?\b",
        r"\byou saved\b",
        r"\bsaved\s*\$?\d+(?:\.\d{2})?\b",
        r"\bcoupon\b",
        r"\bdiscount\b",
        r"\bpromotion\b",
        r"\bpromo\b",
        r"\border totals?\b",
        r"\bitems subtotal\b",
        r"\bsubtotal\b",
        r"\btotal\b",
        r"\btax\b",
        r"\btip\b",
        r"\bservice fee\b",
        r"\bdelivery fee\b",
        r"\bmembership\b",
        r"\bfinal item price\b",
        r"\bpayment method\b",
        r"\bauthorized\b",
        r"\bstatement\b",
        r"^your items$",
        r"^costco$",
    ]
    if any(re.search(pat, n) for pat in summary_patterns):
        return True
    if len(n.split()) >= 8 and any(tok in n for tok in (" your ", " was ", " were ", " authorized ", " statement ")):
        return True
    if n == n.upper() and len(n.split()) <= 4:
        return True
    return False


def clean_order_items(items: List[OrderItem], ai_engine: Optional[AICategoryEngine] = None) -> List[OrderItem]:
    grouped: Dict[str, OrderItem] = {}

    for item in items:
        name = re.sub(r"\s+", " ", item.name).strip()
        if is_summary_item_name(name):
            continue

        quantity_locked = bool(getattr(item, "quantity_locked", False))
        qty_txt = normalize_quantity(item.quantity)
        qty = int(qty_txt) if qty_txt else None
        unit = parse_money_value(item.unit_price)
        total = parse_money_value(item.total_price)
        weight_qty = item.weight_qty
        weight_unit = item.weight_unit.lower() if item.weight_unit else None
        original_total = parse_money_value(item.original_total_price)
        is_weighted = bool(weight_unit)

        if qty is None and unit and total and unit > 0 and not quantity_locked and not is_weighted:
            ratio = total / unit
            rounded = int(round(ratio))
            if rounded >= 1 and abs(ratio - rounded) < 0.08:
                qty = rounded
                if debug_item_match(name):
                    debug_item_log(
                        f"qty inferred from ratio (missing qty) -> name={name!r}, unit={unit}, total={total}, qty={qty}"
                    )
        elif qty is not None and unit and total and unit > 0 and not quantity_locked and not is_weighted:
            # If scraped quantity is too low but line total implies a higher integer quantity,
            # trust the arithmetic (common on Instacart quantity rendering variants).
            ratio = total / unit
            rounded = int(round(ratio))
            if rounded > qty and abs(ratio - rounded) < 0.08:
                if debug_item_match(name):
                    debug_item_log(
                        f"qty upgraded from ratio -> name={name!r}, old_qty={qty}, unit={unit}, total={total}, new_qty={rounded}"
                    )
                qty = rounded

        if qty is not None and unit is not None and not is_weighted:
            expected = unit * qty
            if total is None or total < unit or abs(total - expected) > 0.02:
                total = expected

        key = normalize_for_match(name)
        candidate = OrderItem(
            name=name,
            quantity=str(qty) if qty is not None else None,
            unit_price=format_money(unit),
            total_price=format_money(total),
            weight_qty=weight_qty,
            weight_unit=weight_unit,
            original_total_price=format_money(original_total) if original_total is not None else None,
            quantity_locked=quantity_locked,
        )

        existing = grouped.get(key)
        if existing is None:
            grouped[key] = candidate
            continue

        existing_qty = int(existing.quantity) if existing.quantity else 0
        candidate_qty = int(candidate.quantity) if candidate.quantity else 0

        existing_unit = parse_money_value(existing.unit_price)
        candidate_unit = parse_money_value(candidate.unit_price)
        existing_total = parse_money_value(existing.total_price)
        candidate_total = parse_money_value(candidate.total_price)
        existing_weighted = bool(existing.weight_unit)
        candidate_weighted = bool(candidate.weight_unit)

        def completeness_score(item: OrderItem) -> int:
            score = 0
            if item.quantity:
                score += 3
            if item.unit_price:
                score += 2
            if item.total_price:
                score += 2
            return score

        effective_existing_qty = existing_qty if existing_qty > 0 else 1
        effective_candidate_qty = candidate_qty if candidate_qty > 0 else 1
        existing_total_consistent = (
            existing_total is not None
            and existing_unit is not None
            and abs(existing_total - effective_existing_qty * existing_unit) < 0.02
        )
        candidate_total_consistent = (
            candidate_total is not None
            and candidate_unit is not None
            and abs(candidate_total - effective_candidate_qty * candidate_unit) < 0.02
        )

        choose_candidate = False
        if existing_weighted or candidate_weighted:
            if candidate_weighted and not existing_weighted:
                choose_candidate = True
            elif existing_weighted and not candidate_weighted:
                choose_candidate = False
            elif (
                candidate_total is not None
                and existing_total is not None
                and candidate_unit is not None
                and existing_unit is not None
                and abs(candidate_unit - existing_unit) < 0.001
                and candidate_total < existing_total
            ):
                # Weighted duplicates often include original and discounted totals; keep discounted total.
                choose_candidate = True
            elif (
                candidate.weight_qty is not None
                and (existing.weight_qty is None or candidate.weight_qty > existing.weight_qty)
                and candidate_total is not None
            ):
                choose_candidate = True

        if (
            not choose_candidate
            and candidate_unit is not None
            and existing_unit is not None
            and candidate_total_consistent
            and existing_total_consistent
            and candidate_unit < existing_unit
            and candidate_unit > 0
        ):
            # Discount-first preference: keep lower unit price for same normalized item.
            choose_candidate = True
        elif not choose_candidate and completeness_score(candidate) > completeness_score(existing):
            choose_candidate = True
        elif not choose_candidate and candidate_qty > existing_qty:
            choose_candidate = True
        elif not choose_candidate and (
            candidate_unit is not None
            and existing_unit is not None
            and effective_candidate_qty == effective_existing_qty
            and candidate_total_consistent
            and existing_total_consistent
            and candidate_unit < existing_unit
        ):
            # Prefer discounted rows when duplicate cards contain both current and original price variants.
            choose_candidate = True
        elif not choose_candidate and (
            candidate_total is not None
            and existing_total is not None
            and candidate_qty > 0
            and existing_qty > 0
            and abs(candidate_total - candidate_qty * (candidate_unit or 0.0)) < 0.02
            and abs(existing_total - existing_qty * (existing_unit or 0.0)) > 0.02
        ):
            choose_candidate = True
        elif (
            not choose_candidate
            and not (existing_weighted or candidate_weighted)
            and candidate_total is not None
            and (existing_total is None or candidate_total > existing_total)
        ):
            choose_candidate = True
        elif not choose_candidate and existing.quantity is None and candidate.quantity is not None:
            choose_candidate = True

        if choose_candidate:
            grouped[key] = candidate

    cleaned = list(grouped.values())
    for item in cleaned:
        cat, source = deterministic_category(item.name)
        item.budget_category = canonicalize_category(cat) or DEFAULT_CATEGORY
        item.category_source = source

    if ai_engine is not None:
        pending_by_key: Dict[str, List[OrderItem]] = {}
        batch_items: Dict[str, Dict[str, str]] = {}
        for item in cleaned:
            if item.category_source != 'default':
                continue
            cache_key = normalize_for_match(item.name)
            cached = canonicalize_category(ai_engine.get_cached_category_by_key(cache_key))
            if cached in CATEGORY_ORDER:
                item.budget_category = cached
                item.category_source = 'ai_cache'
            else:
                pending_by_key.setdefault(cache_key, []).append(item)
                if cache_key not in batch_items:
                    batch_items[cache_key] = {
                        'id': cache_key,
                        'sku': '',
                        'desc': item.name,
                    }

        if batch_items:
            batch_results: Dict[str, Optional[str]] = {}
            for batch in ai_engine.chunk_list(list(batch_items.values())):
                batch_results.update(ai_engine.try_category_batch(batch, source_label='Instacart item'))
            for cache_key, item_list in pending_by_key.items():
                ai_category = canonicalize_category(batch_results.get(cache_key))
                if ai_category in CATEGORY_ORDER:
                    ai_engine.set_cached_category_by_key(cache_key, ai_category)
                    for item in item_list:
                        item.budget_category = ai_category
                        item.category_source = 'ai'
                else:
                    ai_engine.set_cached_category_by_key(cache_key, None)

    cleaned.sort(key=lambda x: x.name.lower())
    return cleaned


_BANNED_ITEM_NAME_PATTERNS: Tuple[str, ...] = (
    r"^receipt$",
    r"^for you,?$",
    r"^for you,\s*for a friend$",
    r"^,?\s*for a friend$",
    r"^friends can get .*terms apply\.?$",
    r".*\bterms apply\.?$",
    r"^recycling fee$",
    r"^[\.\d\s]*kg$",
    r"^found\s*\(\d+\)$",
    r"^current price:?$",
    r"^current price:\s*\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?$",
    r"^price:?$",
    r"^price:\s*\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?$",
    r"^\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?\s*(?:-?\s*each|each)?$",
    r"^final item price:?$",
    r"^item price:?$",
    r"^subtotal:?$",
    r"^tax:?$",
    r"^tip:?$",
    r"^total:?$",
    r"^order total:?$",
    r"^charged:?$",
)


def _collect_item_candidates(page: Page) -> List[Dict[str, Any]]:
    return page.evaluate(
        """
        () => {
          const out = [];
          const textSeen = new Set();
          const moneyRe = /\\$\\s?\\d+(?:,\\d{3})*(?:\\.\\d{2})?/;
          const pushText = (txt, source) => {
            const text = (txt || '').trim();
            if (!text || textSeen.has(text)) return;
            textSeen.add(text);
            out.push({ text, source });
          };

          for (const img of Array.from(document.querySelectorAll('main img, section img, article img'))) {
            let n = img.parentElement;
            for (let depth = 0; depth < 6 && n; depth += 1, n = n.parentElement) {
              const txt = (n.innerText || '').trim();
              if (!txt || txt.length < 10 || txt.length > 420) continue;
              if (!moneyRe.test(txt) || !/[A-Za-z]/.test(txt)) continue;
              pushText(txt, 'img-ancestor');
            }
          }

          const selectorGroups = [
            '[data-testid*="item"]',
            '[data-testid*="product"]',
            '[data-testid*="line-item"]',
            '[data-testid*="receipt-item"]',
            '[role="listitem"]',
            'li',
            'article',
            'div'
          ];
          const seen = new Set();
          const nodes = [];
          for (const sel of selectorGroups) {
            for (const n of Array.from(document.querySelectorAll(sel))) {
              if (!n || seen.has(n)) continue;
              seen.add(n);
              nodes.push(n);
            }
          }
          for (const n of nodes) {
            const txt = (n.innerText || '').trim();
            if (!txt || txt.length < 3 || txt.length > 420) continue;
            if (!/\\$\\s?\\d/.test(txt)) continue;
            if (!/[A-Za-z]/.test(txt) && txt.length > 80) continue;
            pushText(txt, 'selector');
          }
          return out.slice(0, 800);
        }
        """
    )


def _normalize_candidate_text(txt: str) -> Tuple[str, List[str]]:
    txt = re.sub(r"([A-Za-z])(\$\s?\d)", r"\1\n\2", txt)
    txt = re.sub(r"(•\s*each)(\s*Quantity\s*:?)", r"\1\n\2", txt, flags=re.I)
    txt = re.sub(r"(•\s*each)\s*(\d{1,2})\s*(\$\s?\d)", r"\1\n\2\n\3", txt, flags=re.I)
    txt = re.sub(r"(\S)(Quantity\s*:?)", r"\1\n\2", txt, flags=re.I)
    txt = re.sub(r"(\S)(Current\s+price\s*:?)", r"\1\n\2", txt, flags=re.I)
    txt = re.sub(r"(\S)(Original\s+price\s*:?)", r"\1\n\2", txt, flags=re.I)
    lines = [ln.strip() for ln in txt.splitlines() if ln.strip()]
    return txt, lines


def _extract_name(lines: List[str]) -> Tuple[Optional[str], int]:
    for idx, ln in enumerate(lines):
        if re.fullmatch(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", ln):
            continue
        if re.fullmatch(r"x\s*\d+", ln, re.I):
            continue
        if re.search(r"\bcurrent\s+price\b", ln, re.I):
            continue
        if any(re.fullmatch(pat, ln, re.I) for pat in _BANNED_ITEM_NAME_PATTERNS):
            continue
        if len(ln) >= 3:
            return ln, idx
    return None, -1


def _sanitize_name(name: str) -> Optional[str]:
    name = re.sub(r"quantity\s*:?\s*\d+\b", " ", name, flags=re.I)
    name = re.sub(r"current\s+price\s*:?\s*\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?", " ", name, flags=re.I)
    name = re.sub(r"original\s+price\s*:?\s*\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?", " ", name, flags=re.I)
    name = re.sub(r"\bfor you,\s*for a friend\b", " ", name, flags=re.I)
    name = re.sub(r"^,?\s*for a friend\b", " ", name, flags=re.I)
    name = re.sub(r"friends can get .*terms apply\.?", " ", name, flags=re.I)
    name = re.sub(r"\bterms apply\.?", " ", name, flags=re.I)
    name = re.sub(r"\bfor you,?\b", " ", name, flags=re.I)
    name = re.sub(r"\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?\s*(?:•\s*)?each\b", " ", name, flags=re.I)
    name = re.sub(r"\$\s*\d+(?:,\d{3})*(?:\.\d{2})?\b", " ", name)
    name = re.sub(r"\$?\s*\d+(?:,\d{3})*(?:\.\d{2})?\s*earned\.?$", " ", name, flags=re.I)
    name = re.sub(r"\s+", " ", name).strip(" -:|")
    if (
        not name
        or re.search(r"\b(?:quantity|current\s+price|original\s+price)\b", name, re.I)
        or re.search(r"\bearned\.?$", name, re.I)
        or re.search(r"\bterms apply\.?$", name, re.I)
        or re.search(r"\brecycling fee\b", name, re.I)
        or re.fullmatch(r"[\.\d\s]*kg", name, re.I) is not None
        or not re.search(r"[A-Za-z]", name)
    ):
        return None
    return name


def _merge_fragment_into_previous_item(items: List[OrderItem], txt: str, lines: List[str]) -> None:
    if not items:
        return
    qty_only_lines = [ln for ln in lines if re.fullmatch(r"\d{1,2}", ln)]
    if qty_only_lines and not any(re.search(r"[A-Za-z]", ln) for ln in lines):
        last = items[-1]
        parsed_qty = max(int(x) for x in qty_only_lines)
        prev_qty = int(last.quantity) if last.quantity and str(last.quantity).isdigit() else 0
        if parsed_qty > prev_qty:
            if debug_item_match(last.name, txt):
                debug_item_log(
                    f"fragment qty merged to previous item -> prev_name={last.name!r}, prev_qty={prev_qty}, new_qty={parsed_qty}, lines={lines!r}"
                )
            last.quantity = str(parsed_qty)
            prev_unit = parse_money_value(last.unit_price)
            if prev_unit is not None:
                last.total_price = format_money(prev_unit * parsed_qty)

    hint_has_current = bool(re.search(r"\bcurrent\s+price\b", txt, re.I))
    money_only_lines = [ln for ln in lines if re.search(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", ln)]
    if hint_has_current or (money_only_lines and len(lines) <= 4):
        vals = [parse_money_value(m) for m in re.findall(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", txt)]
        vals = [v for v in vals if v is not None and v > 0]
        if vals:
            candidate_unit = min(vals)
            last = items[-1]
            prev_unit = parse_money_value(last.unit_price)
            if prev_unit is None or candidate_unit < prev_unit:
                if debug_item_match(last.name, txt):
                    debug_item_log(
                        f"fragment price merged to previous item -> prev_name={last.name!r}, prev_unit={prev_unit}, new_unit={candidate_unit}, lines={lines!r}"
                    )
                last.unit_price = format_money(candidate_unit)
                if last.quantity and str(last.quantity).isdigit():
                    last.total_price = format_money(candidate_unit * int(last.quantity))
                else:
                    last.total_price = format_money(candidate_unit)


def _queue_fragment(items: List[OrderItem], txt: str, lines: List[str], pending_fragments: List[Dict[str, Any]]) -> None:
    fragment_qty: Optional[int] = None
    fragment_vals: List[float] = [
        v for v in (parse_money_value(m) for m in re.findall(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", txt))
        if v is not None and v > 0
    ]
    qty_tokens = [int(ln) for ln in lines if re.fullmatch(r"\d{1,2}", ln)]
    if qty_tokens:
        fragment_qty = max(qty_tokens)
    else:
        m_qty = re.search(r"\b(\d{1,2})\b(?=[^\n$]{0,8}\$\s?\d)", txt)
        if m_qty:
            fragment_qty = int(m_qty.group(1))

    _merge_fragment_into_previous_item(items, txt, lines)
    if fragment_qty is None and not fragment_vals:
        return

    anchor_key = normalize_for_match(items[-1].name) if items else None
    if debug_item_match(txt):
        debug_item_log(
            f"queued fragment -> qty={fragment_qty}, vals={fragment_vals}, anchor={anchor_key!r}, lines={lines!r}"
        )
    pending_fragments.append({"qty": fragment_qty, "vals": fragment_vals, "anchor_key": anchor_key})


def _extract_initial_quantity(txt: str, lines: List[str], name_idx: int, is_weighted_item: bool) -> Tuple[Optional[str], bool]:
    quantity: Optional[str] = None
    quantity_locked = is_weighted_item
    has_weight_quantity = bool(
        re.search(r"\b(?:qty|quantity)\.?\s*[:x]?\s*\d+(?:\.\d+)?\s*(kg|g|lb|lbs|oz)\b", txt, re.I)
        or re.search(r"\b\d+(?:\.\d+)?\s*(kg|g|lb|lbs|oz)\b", txt, re.I)
    )
    if not has_weight_quantity:
        q = re.search(r"\b(?:qty|quantity)\.?\s*[:x]?\s*(\d+)\b", txt, re.I)
        if not q:
            q = re.search(r"\b(\d+)\s*[xX]\b", txt)
        if not q:
            q = re.search(r"\bx\s*(\d+)\b", txt, re.I)
        if q:
            quantity = q.group(1)

    if quantity is None:
        qty_positions: List[Tuple[int, int]] = []
        money_line_idxs = [i for i, ln in enumerate(lines) if re.search(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", ln)]
        for i, ln in enumerate(lines):
            if not re.fullmatch(r"\d{1,2}", ln):
                continue
            if i > 0 and re.search(r"(kg|g|lb|lbs|oz)\b", lines[i - 1], re.I):
                continue
            if i + 1 < len(lines) and re.search(r"(kg|g|lb|lbs|oz)\b", lines[i + 1], re.I):
                continue
            val = int(ln)
            if val < 1 or val > 99:
                continue
            has_text_before = any(re.search(r"[A-Za-z]", x) for x in lines[:i])
            next_money_idxs = [m for m in money_line_idxs if m > i]
            if not has_text_before or not next_money_idxs:
                continue
            score = next_money_idxs[0] - i
            if i <= name_idx:
                score += 5
            qty_positions.append((score, i))
        if qty_positions:
            _, best_idx = sorted(qty_positions, key=lambda x: x[0])[0]
            quantity = lines[best_idx]

    if re.search(r"quantity\s*:\s*1\b", txt, re.I):
        quantity = "1"
        quantity_locked = True
    return quantity, quantity_locked


def _parse_unit_and_total(
    name: str,
    txt: str,
    lines: List[str],
    quantity: Optional[str],
    quantity_locked: bool,
    is_weighted_item: bool,
) -> Tuple[Optional[str], Optional[str], Optional[float], Optional[str], Optional[str], bool, Optional[str]]:
    unit_price: Optional[str] = None
    total_price: Optional[str] = None
    weight_qty: Optional[float] = None
    weight_unit: Optional[str] = None
    original_total_price: Optional[str] = None

    each_price_match = re.search(r"(\$\s?\d+(?:,\d{3})*(?:\.\d{2})?)\s*[•·]?\s*each", txt, re.I)
    qty_total_match = re.search(r"\b(\d{1,2})\b\s*(\$\s?\d+(?:,\d{3})*(?:\.\d{2})?)\b", txt, re.I)
    if each_price_match and qty_total_match:
        each_val = parse_money_value(each_price_match.group(1))
        qty_val = int(qty_total_match.group(1))
        total_val = parse_money_value(qty_total_match.group(2))
        if each_val is not None and total_val is not None and qty_val >= 1 and abs((each_val * qty_val) - total_val) < 0.08:
            quantity = str(qty_val)
            unit_price = format_money(each_val)
            total_price = format_money(total_val)
            if qty_val == 1:
                quantity_locked = True
            if debug_item_match(name, txt):
                debug_item_log(
                    f"direct each/qty/total pattern -> name={name!r}, qty={quantity}, unit={unit_price}, total={total_price}, lines={lines!r}"
                )

    if unit_price is None or total_price is None:
        current_match = re.search(r"current\s+price\s*:?\s*(\$\s?\d+(?:,\d{3})*(?:\.\d{2})?)", txt, re.I)
        original_match = re.search(r"original\s+price\s*:?\s*(\$\s?\d+(?:,\d{3})*(?:\.\d{2})?)", txt, re.I)
        each_match = re.search(r"(\$\s?\d+(?:,\d{3})*(?:\.\d{2})?)\s*[•·]?\s*each", txt, re.I)
        unit_measure_match = re.search(r"(\$\s?\d+(?:,\d{3})*(?:\.\d{2})?)\s*/\s*(kg|g|lb|lbs|oz)\b", txt, re.I)
        original_value = parse_money_value(original_match.group(1)) if original_match else None
        if current_match:
            current_value = parse_money_value(current_match.group(1))
            each_value = parse_money_value(each_match.group(1)) if each_match else None
            unit_measure_value = parse_money_value(unit_measure_match.group(1)) if unit_measure_match else None
            qty_num = int(quantity) if quantity and str(quantity).isdigit() else None
            if current_value is not None:
                if is_weighted_item and unit_measure_value is not None:
                    quantity = None
                    wm = re.search(r"\b(\d+(?:\.\d+)?)\s*(kg|g|lb|lbs|oz)\b", txt, re.I)
                    if wm:
                        try:
                            weight_qty = float(wm.group(1))
                        except ValueError:
                            weight_qty = None
                        weight_unit = wm.group(2).lower()
                    unit_price = format_money(unit_measure_value)
                    total_price = format_money(current_value)
                    if weight_qty and unit_measure_value:
                        computed_original = unit_measure_value * weight_qty
                        if abs(computed_original - current_value) > 0.05:
                            original_total_price = format_money(computed_original)
                    quantity_locked = True
                    if debug_item_match(name, txt):
                        debug_item_log(f"weighted item pricing -> name={name!r}, unit={unit_measure_value}, total={current_value}")
                elif each_value is not None and qty_num is not None and qty_num >= 1 and abs((each_value * qty_num) - current_value) < 0.12:
                    unit_price = format_money(each_value)
                    total_price = format_money(current_value)
                    if debug_item_match(name, txt):
                        debug_item_log(
                            f"current-price interpreted as line total -> name={name!r}, qty={qty_num}, each={each_value}, total={current_value}"
                        )
                else:
                    unit_price = format_money(current_value)
                    total_price = format_money(current_value)
        else:
            money = [m.replace(" ", "") for m in re.findall(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", txt)]
            if money:
                money_values = [parse_money_value(m) for m in money]
                money_values = [v for v in money_values if v is not None]
                weighted_resolved = False
                unit_measure_value = parse_money_value(unit_measure_match.group(1)) if unit_measure_match else None
                if is_weighted_item and unit_measure_value is not None and money_values:
                    weight_match = re.search(r"\b(\d+(?:\.\d+)?)\s*(kg|g|lb|lbs|oz)\b", txt, re.I)
                    weighted_total: Optional[float] = None
                    if weight_match:
                        try:
                            weight_value = float(weight_match.group(1))
                        except ValueError:
                            weight_value = 0.0
                        weight_token = weight_match.group(2).lower()
                        unit_label = unit_measure_match.group(2).lower()
                        if weight_value > 0 and (
                            (unit_label == "kg" and weight_token == "kg")
                            or (unit_label in ("lb", "lbs") and weight_token in ("lb", "lbs"))
                        ):
                            weighted_total = unit_measure_value * weight_value
                            weight_qty = weight_value
                            weight_unit = unit_label
                    quantity = None
                    unit_price = format_money(unit_measure_value)
                    total_price = format_money(weighted_total if weighted_total is not None else max(money_values))
                    if weighted_total is not None:
                        max_val = max(money_values)
                        if abs(max_val - weighted_total) > 0.05:
                            original_total_price = format_money(max_val)
                    quantity_locked = True
                    weighted_resolved = True
                    if debug_item_match(name, txt):
                        debug_item_log(
                            f"weighted fallback pricing -> name={name!r}, unit={unit_price}, total={total_price}, weight_match={weight_match.group(0) if weight_match else None}"
                        )
                if not weighted_resolved:
                    each_line_values: List[float] = []
                    non_each_values: List[float] = []
                    for ln in lines:
                        vals = [parse_money_value(m) for m in re.findall(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", ln)]
                        vals = [v for v in vals if v is not None]
                        if not vals:
                            continue
                        if re.search(r"\beach\b", ln, re.I):
                            each_line_values.extend(vals)
                        else:
                            non_each_values.extend(vals)
                    qty_num = int(quantity) if quantity and str(quantity).isdigit() else None
                    if qty_num and each_line_values and non_each_values:
                        each_val = min(each_line_values)
                        total_val = max(non_each_values)
                        implied_qty = int(round(total_val / each_val)) if each_val > 0 else 0
                        if implied_qty >= 1 and abs((implied_qty * each_val) - total_val) < 0.08:
                            if qty_num < implied_qty:
                                qty_num = implied_qty
                                quantity = str(implied_qty)
                            unit_price = format_money(each_val)
                            total_price = format_money(total_val)
                        elif abs(each_val * qty_num - total_val) < 0.08:
                            unit_price = format_money(each_val)
                            total_price = format_money(total_val)
                        else:
                            current_value = min(non_each_values)
                            unit_price = format_money(current_value)
                            total_price = format_money(current_value * qty_num)
                    elif non_each_values:
                        current_value = min(non_each_values)
                        unit_price = format_money(current_value)
                        total_price = format_money(current_value * int(quantity)) if quantity and str(quantity).isdigit() else format_money(current_value)
                    elif len(money_values) >= 2:
                        current_value = min(money_values)
                        original_value = max(money_values)
                        if quantity and str(quantity).isdigit() and original_value > 0:
                            implied_qty = int(round(original_value / current_value))
                            if implied_qty > 1 and abs((implied_qty * current_value) - original_value) < 0.08:
                                quantity = str(implied_qty)
                        unit_price = format_money(current_value)
                        total_price = format_money(current_value * int(quantity)) if quantity and str(quantity).isdigit() else format_money(current_value)
                        if abs(current_value - original_value) < 0.001:
                            total_price = unit_price
                    else:
                        unit_price = money[0]
                        total_price = money[-1]
                    total_value = parse_money_value(total_price)
                    if original_value is not None and total_value is not None and abs(total_value - original_value) < 0.001:
                        total_price = unit_price

        final_total_value = parse_money_value(total_price)
        if (
            original_total_price is None
            and original_value is not None
            and final_total_value is not None
            and original_value > final_total_value + 0.005
        ):
            original_total_price = format_money(original_value)

    return quantity, unit_price, weight_qty, weight_unit, original_total_price, quantity_locked, total_price


def _apply_pending_fragment_to_item(name: str, quantity: Optional[str], unit_price: Optional[str], total_price: Optional[str], quantity_locked: bool, pending_fragments: List[Dict[str, Any]]) -> Tuple[Optional[str], Optional[str]]:
    current_unit = parse_money_value(unit_price)
    current_qty = int(quantity) if quantity and str(quantity).isdigit() else 0
    if not (pending_fragments and current_unit is not None and current_unit > 0 and not quantity_locked):
        return quantity, total_price

    matched_idx: Optional[int] = None
    matched_qty = current_qty
    matched_total: Optional[float] = parse_money_value(total_price)
    name_key = normalize_for_match(name)
    for idx, frag in enumerate(pending_fragments):
        frag_anchor = frag.get("anchor_key")
        if frag_anchor and frag_anchor != name_key:
            continue
        p_qty = frag.get("qty")
        p_vals = [v for v in (frag.get("vals") or []) if isinstance(v, (int, float)) and v > 0]
        candidate_qty: Optional[int] = None
        candidate_total: Optional[float] = None
        if p_qty:
            if p_vals:
                nearest = min(p_vals, key=lambda v: abs((current_unit * p_qty) - v))
                if abs((current_unit * p_qty) - nearest) < 0.12:
                    candidate_qty = int(p_qty)
                    candidate_total = float(nearest)
                else:
                    candidate_qty = int(p_qty)
                    candidate_total = float(current_unit * p_qty)
            else:
                candidate_qty = int(p_qty)
                candidate_total = float(current_unit * p_qty)
        else:
            for v in p_vals:
                implied = int(round(v / current_unit))
                if implied >= 1 and abs((implied * current_unit) - v) < 0.12:
                    if candidate_qty is None or implied > candidate_qty:
                        candidate_qty = implied
                        candidate_total = float(v)
        if candidate_qty is not None and candidate_qty > matched_qty:
            matched_idx = idx
            matched_qty = candidate_qty
            matched_total = candidate_total

    if matched_idx is not None:
        if debug_item_match(name):
            debug_item_log(f"pending fragment applied -> name={name!r}, qty={matched_qty}, total={matched_total}, unit={current_unit}")
        quantity = str(matched_qty)
        total_price = format_money(matched_total) if matched_total is not None else format_money(current_unit * matched_qty)
        pending_fragments.pop(matched_idx)
    return quantity, total_price


def _final_reconcile_fragments(items: List[OrderItem], pending_fragments: List[Dict[str, Any]]) -> None:
    if not pending_fragments or not items:
        return
    for frag in pending_fragments:
        frag_anchor = frag.get("anchor_key")
        if not frag_anchor:
            continue
        p_qty = frag.get("qty")
        p_vals = [v for v in (frag.get("vals") or []) if isinstance(v, (int, float)) and v > 0]
        best_idx: Optional[int] = None
        best_qty: Optional[int] = None
        best_total: Optional[float] = None
        best_error = 1e9

        for idx, it in enumerate(items):
            if normalize_for_match(it.name) != frag_anchor:
                continue
            unit = parse_money_value(it.unit_price)
            if unit is None or unit <= 0 or it.quantity_locked:
                continue
            current_qty = int(it.quantity) if it.quantity and str(it.quantity).isdigit() else 1
            candidate_qty: Optional[int] = None
            candidate_total: Optional[float] = None
            error = 1e9
            if p_qty:
                candidate_qty = int(p_qty)
                if p_vals:
                    nearest = min(p_vals, key=lambda v: abs((unit * candidate_qty) - v))
                    candidate_total = float(nearest)
                    error = abs((unit * candidate_qty) - candidate_total)
                else:
                    candidate_total = float(unit * candidate_qty)
                    error = 0.0
            else:
                for v in p_vals:
                    implied = int(round(v / unit))
                    if implied < 1:
                        continue
                    err = abs((implied * unit) - v)
                    if err < error:
                        candidate_qty = implied
                        candidate_total = float(v)
                        error = err

            if candidate_qty is None or candidate_qty <= current_qty or error > 0.20:
                continue
            if candidate_qty > (best_qty or 0) or (candidate_qty == (best_qty or 0) and error < best_error):
                best_idx = idx
                best_qty = candidate_qty
                best_total = candidate_total
                best_error = error

        if best_idx is not None and best_qty is not None:
            if debug_item_match(items[best_idx].name):
                debug_item_log(
                    f"final fragment reconciliation -> name={items[best_idx].name!r}, qty={best_qty}, total={best_total}, error={best_error}"
                )
            items[best_idx].quantity = str(best_qty)
            if best_total is not None:
                items[best_idx].total_price = format_money(best_total)
            else:
                unit = parse_money_value(items[best_idx].unit_price)
                if unit is not None:
                    items[best_idx].total_price = format_money(unit * best_qty)


def extract_items_from_dom(page: Page) -> List[OrderItem]:
    items: List[OrderItem] = []
    seen: Set[str] = set()
    pending_fragments: List[Dict[str, Any]] = []

    for row in _collect_item_candidates(page):
        raw_txt = row.get("text", "")
        if debug_item_match(raw_txt):
            debug_item_log(f"raw candidate text: {raw_txt[:280]!r}")
        txt, lines = _normalize_candidate_text(raw_txt)
        if not lines:
            continue

        name, name_idx = _extract_name(lines)
        if not name:
            _queue_fragment(items, txt, lines, pending_fragments)
            continue

        name = _sanitize_name(name) or ""
        if not name or any(re.fullmatch(pat, name, re.I) for pat in _BANNED_ITEM_NAME_PATTERNS):
            continue

        is_weighted_item = bool(
            re.search(r"/\s*(kg|g|lb|lbs|oz)\b", txt, re.I) or re.search(r"\b\d+(?:\.\d+)?\s*(kg|g|lb|lbs|oz)\b", txt, re.I)
        )
        quantity, quantity_locked = _extract_initial_quantity(txt, lines, name_idx, is_weighted_item)
        quantity, unit_price, weight_qty, weight_unit, original_total_price, quantity_locked, total_price = _parse_unit_and_total(
            name, txt, lines, quantity, quantity_locked, is_weighted_item
        )
        quantity, total_price = _apply_pending_fragment_to_item(
            name, quantity, unit_price, total_price, quantity_locked, pending_fragments
        )

        key = f"{name.lower()}|{quantity}|{unit_price}|{total_price}"
        if key in seen:
            continue
        seen.add(key)

        item = OrderItem(
            name=name,
            quantity=quantity,
            unit_price=unit_price,
            total_price=total_price,
            quantity_locked=quantity_locked,
            weight_qty=weight_qty,
            weight_unit=weight_unit,
            original_total_price=original_total_price,
        )
        items.append(item)
        if debug_item_match(name):
            debug_item_log(f"parsed item append -> name={name!r}, qty={quantity}, unit={unit_price}, total={total_price}")

    _final_reconcile_fragments(items, pending_fragments)
    return items[:250]


def wait_for_items_to_render(page: Page, timeout_seconds: float = 20.0) -> List[OrderItem]:
    deadline = time.time() + max(2.0, timeout_seconds)
    best: List[OrderItem] = []
    stable_hits = 0

    while time.time() < deadline:
        # Expand lazy sections where present.
        for label in ("show all", "view all", "see all", "more items", "show more"):
            try:
                btn = page.get_by_text(label, exact=False).first
                if btn.count() and btn.is_visible():
                    btn.click(timeout=800)
                    time.sleep(0.25)
            except Exception:
                pass

        try:
            page.mouse.wheel(0, 1400)
        except Exception:
            pass
        time.sleep(0.45)

        parsed = clean_order_items(extract_items_from_dom(page))
        if len(parsed) > len(best):
            best = parsed
            stable_hits = 0
        elif len(parsed) == len(best) and len(parsed) > 0:
            stable_hits += 1
        else:
            stable_hits = max(0, stable_hits - 1)

        if stable_hits >= 2 and len(best) > 0:
            break

    return best


def reconcile_items_with_body_text(items: List[OrderItem], body_text: str) -> None:
    if not items or not body_text:
        return
    text = re.sub(r"\s+", " ", body_text)

    for item in items:
        unit = parse_money_value(item.unit_price)
        if unit is None or unit <= 0:
            continue
        cur_qty = int(item.quantity) if item.quantity and str(item.quantity).isdigit() else 1
        name_key = normalize_for_match(item.name)
        if not name_key:
            continue

        # Use first 2-3 significant tokens as anchor to find the item region in page text.
        tokens = [t for t in name_key.split() if len(t) >= 3][:3]
        if len(tokens) < 2:
            continue
        anchor = r"\s+".join(re.escape(t) for t in tokens[:2])

        best_qty = cur_qty
        best_total: Optional[float] = parse_money_value(item.total_price)
        for m in re.finditer(anchor, text, flags=re.I):
            start = max(0, m.start() - 40)
            end = min(len(text), m.end() + 220)
            window = text[start:end]
            qtys = [int(x) for x in re.findall(r"\b([1-9]\d?)\b", window)]
            prices = [parse_money_value(p) for p in re.findall(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", window)]
            prices = [p for p in prices if p is not None and p > 0]
            if not qtys or not prices:
                continue

            # Prefer totals that fit unit * qty; choose the highest matching qty.
            for q in sorted(set(qtys), reverse=True):
                if q <= best_qty:
                    continue
                nearest = min(prices, key=lambda v: abs((unit * q) - v))
                if abs((unit * q) - nearest) < 0.12:
                    best_qty = q
                    best_total = nearest
                    break

        if best_qty > cur_qty:
            if debug_item_match(item.name):
                debug_item_log(
                    f"body-text reconciliation upgrade -> name={item.name!r}, old_qty={cur_qty}, new_qty={best_qty}, total={best_total}"
                )
            item.quantity = str(best_qty)
            if best_total is not None:
                item.total_price = format_money(best_total)
            else:
                item.total_price = format_money(unit * best_qty)

        # Additional strict fallback: anchor on the actual item name and infer qty
        # from nearby prices matching unit * N (useful when qty text is not exposed).
        name_tokens = re.findall(r"[A-Za-z0-9%]+", item.name)
        if len(name_tokens) >= 2:
            name_anchor = r"\s+".join(re.escape(t) for t in name_tokens[:6])
            try:
                name_re = re.compile(name_anchor, re.I)
            except re.error:
                name_re = None
            if name_re:
                for m in name_re.finditer(text):
                    window = text[m.start(): min(len(text), m.end() + 240)]
                    window_prices = [
                        parse_money_value(p)
                        for p in re.findall(r"\$\s?\d+(?:,\d{3})*(?:\.\d{2})?", window)
                    ]
                    window_prices = [p for p in window_prices if p is not None and p > 0]
                    if not window_prices:
                        continue
                    cur_qty2 = int(item.quantity) if item.quantity and str(item.quantity).isdigit() else 1
                    best_q = cur_qty2
                    best_t = parse_money_value(item.total_price)
                    for q in range(max(2, cur_qty2 + 1), 10):
                        target = unit * q
                        nearest = min(window_prices, key=lambda v: abs(v - target))
                        if abs(nearest - target) < 0.12 and q > best_q:
                            best_q = q
                            best_t = nearest
                    if best_q > cur_qty2:
                        if debug_item_match(item.name):
                            debug_item_log(
                                f"name-anchor reconciliation upgrade -> name={item.name!r}, old_qty={cur_qty2}, new_qty={best_q}, total={best_t}"
                            )
                        item.quantity = str(best_q)
                        if best_t is not None:
                            item.total_price = format_money(best_t)
                        else:
                            item.total_price = format_money(unit * best_q)
                        break


def extract_order_metadata(
    page: Page,
    order_url: str,
    ai_engine: Optional[AICategoryEngine] = None,
    store_hint: Optional[str] = None,
) -> OrderRecord:
    title = page.title()
    body_text = wait_for_meaningful_body_text(page, timeout_seconds=10.0)
    receipt_text = capture_receipt_text(page, timeout_seconds=8.0)
    receipt_breakdown = parse_receipt_breakdown(f"{body_text}\n{receipt_text}")
    raw_items = wait_for_items_to_render(page, timeout_seconds=20.0)
    items = clean_order_items(raw_items, ai_engine=ai_engine)
    order_id = parse_order_id(title=title, body_text=body_text, order_url=order_url)

    def infer_fields(text: str) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        inferred_total = parse_order_total(text) if text else None
        inferred_store = (
            parse_store_name(text)
            or clean_store_name(store_hint)
        )
        if not inferred_store:
            inferred_store = wait_for_store_name(page, timeout_seconds=12.0)
        inferred_store = clean_store_name(inferred_store)
        inferred_date = parse_order_date(text) if text else None
        return inferred_total, inferred_store, inferred_date

    order_total, store_name, order_date = infer_fields(body_text)
    if receipt_breakdown.get("total_charged"):
        order_total = receipt_breakdown["total_charged"]
    looks_shell = "post checkout" in (title or "").lower()

    if looks_shell and not items and not order_total and not store_name:
        for detail_url in candidate_order_detail_urls(page, order_url):
            try:
                page.goto(detail_url, wait_until="domcontentloaded")
                time.sleep(0.8)
            except Exception:
                continue
            title = page.title()
            body_text = wait_for_meaningful_body_text(page, timeout_seconds=7.0)
            receipt_text = capture_receipt_text(page, timeout_seconds=6.0)
            receipt_breakdown = parse_receipt_breakdown(f"{body_text}\n{receipt_text}")
            raw_items = wait_for_items_to_render(page, timeout_seconds=14.0)
            items = clean_order_items(raw_items, ai_engine=ai_engine)
            order_total, store_name, order_date = infer_fields(body_text)
            if receipt_breakdown.get("total_charged"):
                order_total = receipt_breakdown["total_charged"]
            if items or order_total or store_name or order_date:
                break

    reconcile_items_with_body_text(items, body_text)
    if not store_name:
        store_name = infer_store_name_from_items(items)

    if not order_date:
        order_date = parse_order_date_from_order_id(order_id)
    items_sum = 0.0
    for item in items:
        val = parse_money_value(item.total_price)
        if val is not None:
            items_sum += val
    order_total_num = parse_money_value(order_total)
    if items_sum > 0 and (order_total_num is None or order_total_num < items_sum * 0.5):
        if not receipt_breakdown.get("total_charged"):
            order_total = format_money(items_sum)

    return OrderRecord(
        order_url=order_url,
        order_id=order_id,
        order_date=order_date,
        store_name=store_name,
        order_total=order_total,
        receipt_item_subtotal=receipt_breakdown.get("item_subtotal"),
        receipt_tip=receipt_breakdown.get("tip"),
        receipt_service_fee=receipt_breakdown.get("service_fee"),
        receipt_recycling_fee=receipt_breakdown.get("recycling_fee"),
        receipt_service_fee_tax=receipt_breakdown.get("service_fee_tax"),
        receipt_gst=receipt_breakdown.get("gst"),
        receipt_pst=receipt_breakdown.get("pst"),
        receipt_discount_total=receipt_breakdown.get("discount_total"),
        receipt_total_charged=receipt_breakdown.get("total_charged"),
        item_count=len(items),
        items=items,
        raw_page_title=title,
    )


def autosize_sheet(ws) -> None:
    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        ws.column_dimensions[letter].width = min(max_len + 2, 60)


def save_outputs(orders: List[OrderRecord], out_json: Path, out_csv: Path, out_xlsx: Path) -> None:
    out_json.write_text(
        json.dumps([asdict(o) for o in orders], ensure_ascii=True, indent=2),
        encoding="utf-8",
    )

    with out_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "order_id",
                "order_date",
                "store_name",
                "order_total",
                "order_url",
                "item_name",
                "budget_category",
                "category_source",
                "unit_qty",
                "weight_qty",
                "weight_unit",
                "unit_cost",
                "line_total",
                "original_line_total",
            ],
        )
        writer.writeheader()
        for order in orders:
            for item in order.items:
                unit_qty = derive_unit_qty(item)
                weight_qty = derive_weight_qty(item)
                writer.writerow(
                    {
                        "order_id": order.order_id,
                        "order_date": order.order_date,
                        "store_name": order.store_name,
                        "order_total": order.order_total,
                        "order_url": order.order_url,
                        "item_name": item.name,
                        "budget_category": item.budget_category,
                        "category_source": item.category_source,
                        "unit_qty": format_unit_qty_for_csv(unit_qty),
                        "weight_qty": format_unit_qty_for_csv(weight_qty),
                        "weight_unit": format_weight_unit_for_csv(item.weight_unit),
                        "unit_cost": item.unit_price,
                        "line_total": item.total_price,
                        "original_line_total": item.original_total_price,
                    }
                )

    last_order_id = next((o.order_id for o in reversed(orders) if o.order_id), None)

    wb = Workbook()
    ws_receipts = wb.active
    ws_receipts.title = "Receipts"
    receipt_headers = [
        "order_id",
        "order_date",
        "store_name",
        "receipt_item_subtotal",
        "receipt_discount_total",
        "receipt_tip",
        "receipt_service_fee",
        "receipt_recycling_fee",
        "receipt_service_fee_tax",
        "receipt_gst",
        "receipt_pst",
        "receipt_total_charged",
        "order_url",
    ]
    ws_receipts.append(receipt_headers)
    for row_idx, order in enumerate(orders, start=2):
        receipt_item_subtotal_num = parse_money_value(order.receipt_item_subtotal)
        receipt_discount_total_num = parse_money_value(order.receipt_discount_total)
        receipt_tip_num = parse_money_value(order.receipt_tip)
        receipt_service_fee_num = parse_money_value(order.receipt_service_fee)
        receipt_recycling_fee_num = parse_money_value(order.receipt_recycling_fee)
        receipt_service_fee_tax_num = parse_money_value(order.receipt_service_fee_tax)
        receipt_gst_num = parse_money_value(order.receipt_gst)
        receipt_pst_num = parse_money_value(order.receipt_pst)
        receipt_total_charged_num = parse_money_value(order.receipt_total_charged)
        ws_receipts.append(
            [
                order.order_id,
                order.order_date,
                order.store_name,
                receipt_item_subtotal_num,
                receipt_discount_total_num,
                receipt_tip_num,
                receipt_service_fee_num,
                receipt_recycling_fee_num,
                receipt_service_fee_tax_num,
                receipt_gst_num,
                receipt_pst_num,
                receipt_total_charged_num,
                order.order_url,
            ]
        )
    autosize_sheet(ws_receipts)
    ws_receipts.auto_filter.ref = f"A1:M{max(1, ws_receipts.max_row)}"
    if last_order_id:
        ws_receipts.auto_filter.add_filter_column(0, [str(last_order_id)])
    for row_idx in range(2, ws_receipts.max_row + 1):
        ws_receipts.cell(row=row_idx, column=4).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=5).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=6).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=7).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=8).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=9).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=10).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=11).number_format = "$#,##0.00"
        ws_receipts.cell(row=row_idx, column=12).number_format = "$#,##0.00"

    ws_items = wb.create_sheet("Items")
    item_headers = [
        "order_id",
        "item_name",
        "budget_category",
        "category_source",
        "unit_qty",
        "weight_qty",
        "weight_unit",
        "unit_cost",
        "line_total",
        "original_line_total",
    ]
    ws_items.append(item_headers)
    for order in orders:
        for item in order.items:
            qty_num_raw = derive_unit_qty(item)
            if qty_num_raw is not None and abs(qty_num_raw - round(qty_num_raw)) < 0.001:
                qty_num = int(round(qty_num_raw))
            else:
                qty_num = qty_num_raw
            weight_qty = derive_weight_qty(item)
            weight_unit = item.weight_unit.upper() if item.weight_unit else None
            unit_num = parse_money_value(item.unit_price)
            total_num = parse_money_value(item.total_price)
            original_total_num = parse_money_value(item.original_total_price)
            ws_items.append(
                [
                    order.order_id,
                    item.name,
                    item.budget_category,
                    item.category_source,
                    qty_num,
                    weight_qty,
                    weight_unit,
                    unit_num,
                    total_num,
                    original_total_num,
                ]
            )
    autosize_sheet(ws_items)
    ws_items.auto_filter.ref = f"A1:J{max(1, ws_items.max_row)}"
    if last_order_id:
        ws_items.auto_filter.add_filter_column(0, [str(last_order_id)])
    for row_idx in range(2, ws_items.max_row + 1):
        qty_cell = ws_items.cell(row=row_idx, column=5)
        qty_val = qty_cell.value
        if isinstance(qty_val, int):
            qty_cell.number_format = "0"
        elif isinstance(qty_val, float):
            if abs(qty_val - round(qty_val)) < 0.001:
                qty_cell.value = int(round(qty_val))
                qty_cell.number_format = "0"
            else:
                qty_cell.number_format = "0.###"
        else:
            qty_cell.number_format = "General"
        wqty_cell = ws_items.cell(row=row_idx, column=6)
        if isinstance(wqty_cell.value, (int, float)):
            wqty_cell.number_format = "0.###"
        ws_items.cell(row=row_idx, column=8).number_format = "$#,##0.00"
        ws_items.cell(row=row_idx, column=9).number_format = "$#,##0.00"
        ws_items.cell(row=row_idx, column=10).number_format = "$#,##0.00"
    wb.save(out_xlsx)


def fill_missing_store_names(orders: List[OrderRecord], default_store_name: Optional[str] = None) -> None:
    explicit_fallback = clean_store_name(default_store_name) if default_store_name else None
    if explicit_fallback:
        for order in orders:
            if not (order.store_name or "").strip():
                order.store_name = explicit_fallback
        return

    known = sorted({(o.store_name or "").strip() for o in orders if (o.store_name or "").strip()})
    if len(known) != 1:
        return
    fallback = known[0]
    for order in orders:
        if not (order.store_name or "").strip():
            order.store_name = fallback


def _require_valid_schema_name(schema: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", schema or ""):
        raise ValueError(f"Invalid --db-schema value: {schema!r}")
    return schema


def _dsn_contains_password(dsn: str) -> bool:
    if not dsn:
        return False
    key_value_password = re.search(r"(^|\s)password\s*=", dsn, re.IGNORECASE)
    uri_password = re.search(r"://[^/\s:@]+:[^@\s/]*@", dsn)
    return bool(key_value_password or uri_password)


def _resolve_postgres_dsn(raw_dsn: str) -> str:
    dsn = (raw_dsn or "").strip()
    if _dsn_contains_password(dsn):
        raise ValueError(
            "Unsafe --db-dsn: inline password detected. "
            "Use ~/.pgpass (recommended) or PGPASSWORD/HOME_BUDGET_PGPASSWORD env vars."
        )
    return dsn or "dbname=home_budget"


def _connect_postgres(dsn: str):
    dsn = _resolve_postgres_dsn(dsn)
    hb_pg_password = os.environ.get("HOME_BUDGET_PGPASSWORD", "").strip()
    if hb_pg_password and not os.environ.get("PGPASSWORD"):
        os.environ["PGPASSWORD"] = hb_pg_password
    try:
        import psycopg  # type: ignore

        conn = psycopg.connect(dsn)
        return conn
    except Exception:
        try:
            import psycopg2  # type: ignore

            conn = psycopg2.connect(dsn)
            return conn
        except Exception as e:
            raise RuntimeError(
                "Could not connect to Postgres. Install psycopg (or psycopg2-binary) and verify --db-dsn. "
                "For auth, prefer ~/.pgpass or libpq env vars."
            ) from e


def write_orders_to_postgres(orders: List[OrderRecord], dsn: str, schema: str) -> Tuple[int, int]:
    schema = _require_valid_schema_name(schema)
    conn = _connect_postgres(dsn)
    inserted_orders = 0
    inserted_items = 0

    order_sql = f"""
        INSERT INTO {schema}.expenses (
            source, order_id, order_date, store_name, order_url,
            receipt_item_subtotal, receipt_discount_total, receipt_tip,
            receipt_service_fee, receipt_recycling_fee, receipt_service_fee_tax,
            receipt_gst, receipt_pst, receipt_total_charged,
            expense_total, raw_page_title, raw_payload
        ) VALUES (
            %s, %s, %s, %s, %s,
            %s, %s, %s,
            %s, %s, %s,
            %s, %s, %s,
            %s, %s, %s::jsonb
        )
        ON CONFLICT (source, order_id)
        DO UPDATE SET
            order_date = EXCLUDED.order_date,
            store_name = EXCLUDED.store_name,
            order_url = EXCLUDED.order_url,
            receipt_item_subtotal = EXCLUDED.receipt_item_subtotal,
            receipt_discount_total = EXCLUDED.receipt_discount_total,
            receipt_tip = EXCLUDED.receipt_tip,
            receipt_service_fee = EXCLUDED.receipt_service_fee,
            receipt_recycling_fee = EXCLUDED.receipt_recycling_fee,
            receipt_service_fee_tax = EXCLUDED.receipt_service_fee_tax,
            receipt_gst = EXCLUDED.receipt_gst,
            receipt_pst = EXCLUDED.receipt_pst,
            receipt_total_charged = EXCLUDED.receipt_total_charged,
            expense_total = EXCLUDED.expense_total,
            raw_page_title = EXCLUDED.raw_page_title,
            raw_payload = EXCLUDED.raw_payload
        RETURNING id
    """
    delete_items_sql = f"DELETE FROM {schema}.expense_items WHERE expense_pk = %s"
    item_sql = f"""
        INSERT INTO {schema}.expense_items (
            expense_pk, item_name, budget_category, category_source,
            unit_qty, unit_cost, weight_qty, weight_unit, line_total, original_line_total
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    """

    try:
        conn.autocommit = False
        with conn.cursor() as cur:
            for order in orders:
                if not order.order_id:
                    print(f"Warning: skipping DB write for order without order_id: url={order.order_url}", flush=True)
                    continue
                order_date_value = None
                if order.order_date:
                    text = order.order_date.strip()
                    for fmt in ("%Y-%m-%d", "%b %d, %Y", "%B %d, %Y", "%m/%d/%Y"):
                        try:
                            order_date_value = datetime.strptime(text, fmt).date()
                            break
                        except ValueError:
                            continue

                cur.execute(
                    order_sql,
                    (
                        "instacart",
                        order.order_id,
                        order_date_value,
                        order.store_name,
                        order.order_url,
                        parse_money_value(order.receipt_item_subtotal),
                        parse_money_value(order.receipt_discount_total),
                        parse_money_value(order.receipt_tip),
                        parse_money_value(order.receipt_service_fee),
                        parse_money_value(order.receipt_recycling_fee),
                        parse_money_value(order.receipt_service_fee_tax),
                        parse_money_value(order.receipt_gst),
                        parse_money_value(order.receipt_pst),
                        parse_money_value(order.receipt_total_charged),
                        parse_money_value(order.order_total),
                        order.raw_page_title,
                        json.dumps(asdict(order), ensure_ascii=True),
                    ),
                )
                row = cur.fetchone()
                if row is None:
                    continue
                expense_pk = int(row[0])
                inserted_orders += 1

                cur.execute(delete_items_sql, (expense_pk,))
                for item in order.items:
                    cur.execute(
                        item_sql,
                        (
                            expense_pk,
                            item.name,
                            item.budget_category,
                            item.category_source,
                            derive_unit_qty(item),
                            parse_money_value(item.unit_price),
                            derive_weight_qty(item),
                            item.weight_unit.lower() if item.weight_unit else None,
                            parse_money_value(item.total_price),
                            parse_money_value(item.original_total_price),
                        ),
                    )
                    inserted_items += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return inserted_orders, inserted_items


def main() -> None:
    args = parse_args()

    profile_dir = Path(args.profile_dir).expanduser().resolve()
    out_json = Path(args.out_json).expanduser().resolve()
    out_csv = Path(args.out_csv).expanduser().resolve()
    out_xlsx = Path(args.out_xlsx).expanduser().resolve()
    ai_engine = AICategoryEngine(AI_SUGGESTIONS_PATH)
    target_month = args.month.strip()
    if target_month:
        if not re.fullmatch(r"\d{4}-\d{2}", target_month):
            raise ValueError(f"Invalid --month format: {target_month} (expected YYYY-MM)")

    profile_dir.mkdir(parents=True, exist_ok=True)
    ai_engine.load_suggestions()

    pw = None
    context: Optional[BrowserContext] = None
    try:
        pw, context = launch_context(profile_dir=profile_dir, headful=args.headful, timeout_ms=args.timeout_ms)
        page = get_or_create_page(context)

        open_orders_page(page)

        if args.interactive_login and not looks_logged_in(page):
            print("Manual login required. Complete login/2FA in the opened browser window.", flush=True)
            wait_for_manual_login(page)
            open_orders_page(page)

        if not looks_logged_in(page):
            raise RuntimeError(
                "Not logged in. Re-run with --headful --interactive-login and complete login in the browser window."
            )

        incremental_scroll(page, args.scroll_rounds)
        order_links = extract_order_links(page)
        order_store_hints = extract_order_store_hints(page)

        if not order_links:
            raise RuntimeError(
                "No order links discovered. Open your orders page in the same browser/profile once, then re-run."
            )

        if args.max_orders > 0:
            order_links = order_links[: args.max_orders]

        print(f"Discovered {len(order_links)} order links", flush=True)

        orders: List[OrderRecord] = []
        all_scraped_orders: List[OrderRecord] = []
        for idx, link in enumerate(order_links, start=1):
            print(f"[{idx}/{len(order_links)}] {link}", flush=True)
            try:
                page.goto(link, wait_until="domcontentloaded")
                time.sleep(0.8)
                order = extract_order_metadata(
                    page,
                    link,
                    ai_engine=ai_engine,
                    store_hint=order_store_hints.get(link),
                )
                print(
                    f'  Parsed item_count={order.item_count}, '
                    f'order_total={order.order_total or "n/a"}, '
                    f'title={order.raw_page_title or "n/a"}',
                    flush=True,
                )
                all_scraped_orders.append(order)
                if month_matches(order.order_date, target_month):
                    orders.append(order)
            except Exception as e:
                print(f"Warning: failed to parse order page {link}: {type(e).__name__}: {e}", flush=True)

        if target_month and not orders and all_scraped_orders:
            print(
                f'Warning: no orders matched --month {target_month}; '
                'keeping all scraped orders to avoid empty output.',
                flush=True,
            )
            orders = all_scraped_orders

        fill_missing_store_names(orders, default_store_name=args.default_store_name)
        try:
            save_outputs(orders, out_json=out_json, out_csv=out_csv, out_xlsx=out_xlsx)
        except PermissionError as e:
            raise RuntimeError(
                f"Could not write output file: {out_xlsx}. "
                "Close the workbook if it is open in Excel and run again."
            ) from e
        if args.write_db:
            written_orders, written_items = write_orders_to_postgres(
                orders,
                dsn=args.db_dsn,
                schema=args.db_schema,
            )
            print(
                f"Wrote Postgres: schema={args.db_schema}, expenses={written_orders}, expense_items={written_items}",
                flush=True,
            )
        ai_engine.save_suggestions()
        print(f"Wrote JSON: {out_json}")
        print(f"Wrote CSV : {out_csv}")
        print(f"Wrote XLSX: {out_xlsx}")
        print(
            'AI Status: '
            f'enabled={ai_engine.enabled}, '
            f'disabled_runtime={ai_engine.runtime_disabled}, '
            f'max_calls_per_run={ai_engine.max_calls_per_run}, '
            f'calls={ai_engine.stats["calls"]}, '
            f'batch_items={ai_engine.stats["batch_items"]}, '
            f'ai_cache_hits={ai_engine.stats["ai_cache_hits"]}, '
            f'skipped_call_budget={ai_engine.stats["skipped_call_budget"]}, '
            f'success={ai_engine.stats["success"]}, '
            f'api_errors={ai_engine.stats["api_errors"]}, '
            f'no_result={ai_engine.stats["no_result"]}.'
        )
        print(f"Orders parsed: {len(orders)}")

    finally:
        try:
            if context is not None:
                context.close()
        finally:
            if pw is not None:
                pw.stop()


if __name__ == "__main__":
    main()
